"""Shared daily report rows, success coloring and target-based estimates."""
from datetime import datetime,timedelta,timezone
from pathlib import Path
import json
import ledger
from report_paths import profile_identity,report_path

CN_TZ=timezone(timedelta(hours=8))
GREEN='E2F0D9'
_workbook_cache={}


def rows(db,profile,at):
    if isinstance(at,str): at=datetime.fromisoformat(at)
    start=at.astimezone(CN_TZ).replace(hour=0,minute=0,second=0,microsecond=0)
    exists=db.execute("SELECT 1 FROM sqlite_master WHERE name='assessments'").fetchone()
    if not exists: return []
    data=db.execute('''SELECT a.*,r.profile_path,c.status AS actual_send_status FROM assessments a
        JOIN runs r ON r.run_id=a.run_id LEFT JOIN candidates c ON c.candidate_hash=a.candidate_hash
        WHERE julianday(a.started_at)>=julianday(?) AND julianday(a.started_at)<julianday(?)
        ORDER BY a.started_at,a.candidate_name''',(start.isoformat(),(start+timedelta(days=1)).isoformat())).fetchall()
    identity=profile_identity(profile)
    return [row for row in data if profile_identity(row['profile_path'])==identity]


def confirmed_sent(row):
    return row['actual_send_status']=='sent'


def projected_counts(db,profile,at):
    data=rows(db,profile,at)
    return dict(viewed=len(data),sent=sum(confirmed_sent(row) for row in data),pending=sum(bool(not row['report_written'] or (row['actual_send_status'] and row['send_status']!=row['actual_send_status'])) for row in data))


def workbook_counts(path):
    from openpyxl import load_workbook
    path=Path(path).resolve()
    before=path.stat();key=(before.st_mtime_ns,before.st_size,before.st_ino)
    cached=_workbook_cache.get(str(path))
    if cached and cached[0]==key:return dict(cached[1])
    wb=load_workbook(path,read_only=False,data_only=True)
    try:
        ws=wb['候选人简报']; viewed=sent=0
        for row in ws.iter_rows(min_row=2):
            if len(row)<2 or not row[1].value: continue
            viewed+=1
            if row[1].fill.fgColor.type=='rgb' and row[1].fill.fgColor.rgb[-6:].upper()==GREEN:
                sent+=1
        result=dict(viewed=viewed,sent=sent)
        after=path.stat()
        if key==(after.st_mtime_ns,after.st_size,after.st_ino):_workbook_cache[str(path)]=(key,result)
        return dict(result)
    finally: wb.close()


def daily_counts(db,run,root,at=None):
    at=at or ledger.now()
    expected=projected_counts(db,run['profile_path'],at)
    path=report_path(Path(root),at.isoformat(),run['profile_path'])
    actual=dict(viewed=0,sent=0)
    error=None
    if path.is_file():
        try: actual=workbook_counts(path)
        except (OSError,ValueError,KeyError) as exc: error=str(exc)
    pending=bool(expected['pending'] or actual['viewed']!=expected['viewed'] or actual['sent']!=expected['sent'] or error)
    return dict(daily_viewed=actual['viewed'],daily_sent=actual['sent'],report_pending=pending,report_path=str(path))


def estimate(run,viewed,sent,average_seconds,elapsed):
    candidates=[]
    if average_seconds and viewed>0:
        candidates.append((max(0,run['max_views']-viewed)*average_seconds,'查看目标'))
    if sent>0 and elapsed>0:
        candidates.append((max(0,run['target']-sent)*elapsed/sent,'招呼目标'))
    if not candidates: return None,'估算中'
    seconds,basis=min(candidates)
    import math
    return math.ceil(seconds/60),basis
