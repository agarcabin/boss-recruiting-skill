"""Offline behavioral acceptance: no live recruiting or desktop actions."""
import contextlib
import io
import sys
import json
import tempfile
import threading
import time
import unittest
from argparse import Namespace
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch,Mock
from PIL import Image
from openpyxl import load_workbook
import ledger
import candidate_log
import recruiting_rules
import run_safety
import safe_action
import stop_controller
import project_registry
import managed_helper
from test_ledger import PROFILE

def call(fn,db,**kw):
    out=io.StringIO()
    with contextlib.redirect_stdout(out): fn(db,Namespace(**kw))
    return json.loads(out.getvalue())

class RecruitingV2(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.profile=self.root/'profile.md'
        self.profile.write_text(PROFILE,encoding='utf-8')
        self.patcher=patch.object(ledger,'PROFILE_ROOT',self.root);self.patcher.start();self.addCleanup(self.patcher.stop)
        self.db=ledger.connect(self.root/'ledger.db');self.addCleanup(self.db.close);candidate_log.table(self.db)
        self.run=call(ledger.start,self.db,profile=str(self.profile),target=150,duration_min=240,min_score=70)['run_id']
        self.rules=json.loads(recruiting_rules.CONFIG.read_text(encoding='utf-8'))
        self.candidate='https://example.test/v2/one'

    def policy(self,name):
        self.rules.update(policy=name,mode='card' if name=='worker' else 'resume',pdf_required=name!='meikang')
        self.db.execute('UPDATE runs SET evaluation_mode=?,rule_snapshot=? WHERE run_id=?',(self.rules['mode'],json.dumps(self.rules),self.run))

    def card(self,advantage='',expectation='普工'):
        self.policy('worker')
        started=call(ledger.candidate_begin,self.db,run=self.run,candidate_id=self.candidate)['candidate_started_at']
        fields=dict(advantage=advantage,expectation=expectation,already_contacted=False)
        image=self.root/'card.png';Image.new('RGB',(80,80),'white').save(image)
        observation=self.root/'card.uia.json'
        observation.write_text(json.dumps(dict(name='测试人',card=fields,observed_at=started,bounds=dict(x=0,y=0,width=80,height=80),document_url='https://www.zhipin.com/web/chat/recommend')),encoding='utf-8')
        return dict(candidate_id=self.candidate,candidate_name='测试人',started_at=started,electronic_resume_enabled=False,
                    overall_score=0,hard_requirements=False,decision='greet',pros='列表命中',cons='其余未知',
                    card=fields,card_evidence=dict(image=str(image),uia_json=str(observation)))

    def record(self,payload):
        return call(lambda db,args:candidate_log.record(db,args,payload),self.db,run=self.run,output_root=self.root)

    def reserve(self): return call(ledger.reserve,self.db,run=self.run,candidate_id=self.candidate)

    def test_worker_expectation_only_bypasses_score_gate_and_resume_requirements(self):
        self.record(self.card())
        row=self.db.execute('SELECT * FROM assessments').fetchone()
        self.assertEqual(row['overall_score'],40);self.assertEqual(row['pdf_status'],'not_applicable')
        self.assertEqual(self.reserve()['state'],'ready')
        self.assertTrue(Path(row['card_image_path']).is_file())
        self.assertIsNone(row['resume_image_path'])

    def test_repeated_initialization_sync_preserve_card_not_applicable(self):
        self.record(self.card('做过组装'))
        reservation=self.reserve()
        call(ledger.resolve,self.db,token=reservation['token'],result='sent',min_gap=80,max_gap=80)
        for _ in range(3):
            candidate_log.table(self.db)
            call(candidate_log.sync,self.db,run=self.run,candidate_id=self.candidate,output_root=self.root)
        row=self.db.execute('SELECT * FROM assessments').fetchone()
        self.assertEqual((row['pdf_status'],row['resume_archive_status']),('not_applicable','not_applicable'))
        for index in (2,3):
            attempt=run_safety.begin_message(self.db,self.run,self.candidate,index)
            run_safety.resolve_message(self.db,attempt['token'],'sent')
        stamp=ledger.iso(ledger.now())
        call(candidate_log.finalize_timing,self.db,run=self.run,candidate_id=self.candidate,ended_at=stamp,output_root=self.root)
        self.assertEqual(call(ledger.candidate_finish,self.db,run=self.run,candidate_id=self.candidate,ended_at=stamp)['state'],'finished_candidate')

    def test_card_image_identity_mismatch_cannot_be_recorded(self):
        payload=self.card();payload['candidate_name']='另一个人'
        with self.assertRaisesRegex(ValueError,'identity_or_fields'): self.record(payload)
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM assessments').fetchone()[0],0)

    def test_keyword_negation_advert_and_positive_after_but(self):
        self.policy('worker')
        for text in ('没做过组装','没有电子厂经验','招聘广告：有组装经验','不考虑普工'):
            result=recruiting_rules.evaluate(self.rules,dict(card=dict(advantage=text,expectation=text)),70)
            self.assertFalse(result['eligible'],text)
        result=recruiting_rules.evaluate(self.rules,dict(card=dict(advantage='没做过组装但有流水线经验',expectation='')),70)
        self.assertTrue(result['eligible']);self.assertEqual(result['score'],60)

    def test_both_worker_groups_score_100_and_contacted_blocks(self):
        self.policy('worker')
        payload=dict(card=dict(advantage='做过组装，电子厂工作经验',expectation='包装工'))
        self.assertEqual(recruiting_rules.evaluate(self.rules,payload,99)['score'],100)
        payload['card']['already_contacted']=True
        self.assertFalse(recruiting_rules.evaluate(self.rules,payload,0)['eligible'])

    def test_agv_unknown_allowed_failed_blocks_certificate_only_does_not_trigger(self):
        self.policy('agv_advanced')
        payload=dict(keyword_text='曾负责海康AGV项目实施',requirement_states={})
        result=recruiting_rules.evaluate(self.rules,payload,70)
        self.assertTrue(result['eligible']);self.assertEqual(result['score'],0)
        payload['requirement_states']={'travel':'failed'}
        self.assertFalse(recruiting_rules.evaluate(self.rules,payload,0)['eligible'])
        payload=dict(keyword_text='持有电工证，已离职',requirement_states=dict(certificate='met',left_job='met'))
        result=recruiting_rules.evaluate(self.rules,payload,70)
        self.assertFalse(result['eligible']);self.assertEqual(result['score'],5)
        self.assertFalse(recruiting_rules.evaluate(self.rules,dict(keyword_text='期望AGV实施工作'),70)['eligible'])

    def test_approved_advanced_accepts_brand_and_agv_from_separate_jobs_at_60(self):
        self.policy('agv_advanced')
        payload=dict(keyword_text='曾在海康做监控安装。后来负责其他品牌AGV项目实施。',
                     requirement_states=dict(brand_work='met',agv='met'),generic_agv_experience=True)
        result=recruiting_rules.evaluate(self.rules,payload,70)
        self.assertTrue(result['eligible']);self.assertEqual(result['score'],60)
        self.assertEqual(result['trigger'],'keyword')
        self.assertEqual(result['requirements']['brand_work'],'met')
        self.assertEqual(result['requirements']['travel'],'unknown')
        for requirement in ('brand_work','agv','independent','vendor','team','travel'):
            with self.subTest(requirement=requirement):
                rejected=dict(payload,requirement_states={**payload['requirement_states'],requirement:'failed'})
                self.assertFalse(recruiting_rules.evaluate(self.rules,rejected,70)['eligible'])

    def test_approved_advanced_unknown_is_not_met_and_generic_bonus_is_removed(self):
        self.policy('agv_advanced')
        payload=dict(keyword_text='负责AGV项目实施',generic_agv_experience=True,
                     requirement_states=dict(certificate='met',left_job='met'))
        result=recruiting_rules.evaluate(self.rules,payload,70)
        self.assertTrue(result['eligible']);self.assertEqual(result['score'],5)
        self.assertEqual(result['requirements']['brand_work'],'unknown')
        self.assertEqual(result['requirements']['agv'],'unknown')
        states={k:'met' for k in ('brand_work','agv','independent','vendor','team','travel','certificate','left_job')}
        full=recruiting_rules.evaluate(self.rules,dict(keyword_text='',requirement_states=states),70)
        self.assertEqual(full['score'],100);self.assertEqual(full['trigger'],'score')

    def test_approved_basic_accepts_either_experience_despite_failed_other_at_low_score(self):
        self.policy('agv_basic')
        cases=[('agv','vendor','曾负责AGV项目实施，没有乙方实施经历',65),
               ('vendor','agv','没有AGV经验。曾在乙方负责项目实施交付',25)]
        for met,failed,text,score in cases:
            with self.subTest(met=met):
                result=recruiting_rules.evaluate(self.rules,dict(keyword_text=text,
                    requirement_states={met:'met',failed:'failed'}),70)
                self.assertTrue(result['eligible']);self.assertEqual(result['trigger'],'keyword')
                self.assertEqual(result['score'],score)
        result=recruiting_rules.evaluate(self.rules,dict(keyword_text='负责AGV与乙方实施项目',
                    requirement_states=dict(agv='failed',vendor='failed')),0)
        self.assertFalse(result['eligible'])

    def test_approved_basic_vendor_negation_advert_intent_and_failed_alternative_do_not_trigger(self):
        self.policy('agv_basic')
        for text in ('没有乙方实施经历','乙方公司未做过实施','招聘广告：乙方实施经验',
                     '希望从事乙方实施工作','乙方销售工作','海康华睿','持有电工证，已离职'):
            with self.subTest(text=text):
                self.assertFalse(recruiting_rules.evaluate(self.rules,dict(keyword_text=text),70)['eligible'])
        result=recruiting_rules.evaluate(self.rules,dict(keyword_text='负责AGV项目实施',
                    requirement_states=dict(agv='failed',vendor='unknown')),70)
        self.assertFalse(result['eligible'])

    def test_version_2_snapshot_keeps_original_weights_and_and_gate(self):
        old=dict(self.rules,version=2,policy='agv_basic',agv_basic_weights=dict(agv=55,vendor=35,certificate=5,left_job=5))
        payload=dict(keyword_text='曾负责AGV项目实施',requirement_states=dict(agv='met',vendor='failed'))
        result=recruiting_rules.evaluate(old,payload,70)
        self.assertFalse(result['eligible']);self.assertEqual(result['score'],55)
        old.update(policy='agv_advanced',agv_advanced_weights=dict(agv=30,independent=20,vendor=20,team=15,travel=5,certificate=5,left_job=5))
        result=recruiting_rules.evaluate(old,dict(keyword_text='',generic_agv_experience=True,
                    requirement_states=dict(certificate='met',left_job='met')),70)
        self.assertEqual(result['score'],30)

    def test_agv_vendor_only_ocr_is_shared_between_record_and_reserve(self):
        self.policy('agv_basic')
        started=call(ledger.candidate_begin,self.db,run=self.run,candidate_id=self.candidate)['candidate_started_at']
        image=self.root/'vendor-resume.png';Image.new('RGB',(80,80),'white').save(image)
        text=self.root/'vendor-ocr.txt';text.write_text('没有AGV经验。曾在乙方负责项目实施交付。'*20,encoding='utf-8')
        ocr=self.root/'vendor-ocr.json';ocr.write_text('{"state":"ok"}',encoding='utf-8')
        payload=dict(candidate_id=self.candidate,candidate_name='乙方实施测试人',started_at=started,electronic_resume_enabled=True,
            overall_score=0,hard_requirements=False,decision='greet',pros='乙方实施经历',cons='无AGV经历',
            electronic=dict(score=0,evidence='乙方实施经历'),resume_evidence=dict(image=str(image),ocr_text=str(text),ocr_json=str(ocr)),
            requirement_states=dict(agv='failed',vendor='met'),
            coworker_progress=dict(state='empty',evidence='offline UIA fixture explicitly empty'))
        self.record(payload)
        row=self.db.execute('SELECT * FROM assessments').fetchone()
        self.assertEqual(row['overall_score'],25)
        self.assertEqual(json.loads(row['eligibility_json'])['trigger'],'keyword')
        self.assertEqual(self.reserve()['state'],'ready')

    def test_meikang_keyword_low_score_and_score_without_keyword(self):
        self.policy('meikang')
        payload=dict(keyword_text='有医疗器械经验，负责海外销售',overall_score=40,hard_requirements=True)
        self.assertEqual(recruiting_rules.evaluate(self.rules,payload,70)['trigger'],'keyword')
        payload.update(keyword_text='',overall_score=80)
        self.assertEqual(recruiting_rules.evaluate(self.rules,payload,70)['trigger'],'score')
        payload['hard_requirements']=False
        self.assertFalse(recruiting_rules.evaluate(self.rules,payload,70)['eligible'])

    def test_agv_score_only_and_any_certificate_non_cumulative(self):
        self.policy('agv_basic')
        payload=dict(requirement_states=dict(agv='met',vendor='met',certificate='met',left_job='met'),keyword_text='')
        result=recruiting_rules.evaluate(self.rules,payload,70)
        self.assertEqual(result['score'],100);self.assertEqual(result['trigger'],'score')

    def test_changed_display_name_preserves_frozen_identity(self):
        before=ledger.run_candidate_hash(self.db,self.run,self.candidate)
        self.profile.write_text(PROFILE.replace('position: Sales','position: Sales（普工）'),encoding='utf-8')
        self.assertEqual(ledger.run_candidate_hash(self.db,self.run,self.candidate),before)
        meta=dict(candidate_namespace='Test Company-Sales',evaluation_mode='card',trigger_policy='worker',display_order=3)
        (self.root/'project.json').write_text(json.dumps(meta),encoding='utf-8')
        project_registry.register_project(self.profile)
        exported=self.root/'export.zip';project_registry.export_project(self.profile,exported)
        imported=project_registry.import_project(exported,self.root/'imported')
        found=list((self.root/'imported').rglob('project.json'))[0]
        self.assertEqual(json.loads(found.read_text(encoding='utf-8'))['candidate_namespace'],'Test Company-Sales')

    def test_300th_can_complete_but_301st_cannot_begin(self):
        at=ledger.iso(ledger.now())
        for index in range(1,300): self.db.execute('INSERT INTO candidate_progress VALUES (?,?,?,?,?,?)',(self.run,'seen-'+str(index),index,at,None,at))
        self.assertEqual(call(ledger.candidate_begin,self.db,run=self.run,candidate_id=self.candidate)['state'],'started_candidate')
        gate=call(ledger.gate,self.db,run=self.run)
        self.assertTrue(gate['may_continue']);self.assertFalse(gate['may_start_candidate'])
        self.assertEqual(call(ledger.candidate_begin,self.db,run=self.run,candidate_id=self.candidate+'2')['reason'],'run_view_limit_reached')

    def test_rolling_window_is_24_hours(self):
        at=ledger.now()
        for index,hours in enumerate((1.5,23.99,24.01)):
            stamp=ledger.iso(at-timedelta(hours=hours))
            self.db.execute('INSERT INTO candidates VALUES (?,?,?,?,?,?,?,?)',(str(index),'test','older','sent',None,stamp,stamp,stamp))
        self.assertEqual(ledger.rolling_count(self.db,at),2)

    def test_150th_first_message_can_finish_followups_but_no_new_candidate(self):
        self.record(self.card());reservation=self.reserve()
        at=ledger.iso(ledger.now())
        for index in range(149):
            self.db.execute('INSERT INTO candidates VALUES (?,?,?,?,?,?,?,?)',('prior-'+str(index),'test',self.run,'sent',None,at,at,at))
        call(ledger.resolve,self.db,token=reservation['token'],result='sent',min_gap=80,max_gap=80)
        gate=call(ledger.gate,self.db,run=self.run)
        self.assertEqual(gate['state'],'draining');self.assertTrue(gate['may_continue']);self.assertFalse(gate['may_start_candidate'])
        self.assertEqual(call(ledger.candidate_finish,self.db,run=self.run,candidate_id=self.candidate)['reason'],'greeting_messages_incomplete')
        second=run_safety.begin_message(self.db,self.run,self.candidate,2)
        run_safety.resolve_message(self.db,second['token'],'sent')
        third=run_safety.begin_message(self.db,self.run,self.candidate,3)
        run_safety.resolve_message(self.db,third['token'],'sent')
        self.assertEqual(ledger.rolling_count(self.db,ledger.now()),150)

    def test_resume_keyword_is_shared_between_record_and_reserve(self):
        self.policy('meikang')
        started=call(ledger.candidate_begin,self.db,run=self.run,candidate_id=self.candidate)['candidate_started_at']
        image=self.root/'resume.png';Image.new('RGB',(80,80),'white').save(image)
        text=self.root/'ocr.txt';text.write_text('有医疗器械经验，负责海外销售。'*20,encoding='utf-8')
        ocr=self.root/'ocr.json';ocr.write_text('{"state":"ok"}',encoding='utf-8')
        payload=dict(candidate_id=self.candidate,candidate_name='测试人',started_at=started,electronic_resume_enabled=True,
            overall_score=40,hard_requirements=True,decision='greet',pros='行业及海外经历',cons='分数较低',
            electronic=dict(score=40,evidence='医疗器械海外销售'),resume_evidence=dict(image=str(image),ocr_text=str(text),ocr_json=str(ocr)),
            coworker_progress=dict(state='empty',evidence='offline UIA fixture explicitly empty'))
        self.record(payload)
        self.assertEqual(self.reserve()['state'],'ready')

    def test_first_success_survives_later_unknown_and_blocks_other_runs(self):
        self.record(self.card());reservation=self.reserve()
        call(ledger.resolve,self.db,token=reservation['token'],result='sent',min_gap=80,max_gap=80)
        second=run_safety.begin_message(self.db,self.run,self.candidate,2)
        run_safety.resolve_message(self.db,second['token'],'unknown')
        self.assertEqual(ledger.rolling_count(self.db,ledger.now()),1)
        self.assertEqual(ledger.unresolved(self.db),1)
        self.assertEqual(call(ledger.start,self.db,profile=str(self.profile),target=150,duration_min=240,min_score=70)['reason'],'unresolved_prior_send')
        with self.assertRaisesRegex(ValueError,'manual'): run_safety.resolve_message(self.db,second['token'],'sent')
        run_safety.resolve_message(self.db,second['token'],'not-sent',manual=True)
        self.assertEqual(ledger.rolling_count(self.db,ledger.now()),1)

    def test_stop_fences_late_resolve_and_finalizes_once(self):
        self.record(self.card());reservation=self.reserve()
        run_safety.request_stop(self.db,self.run,'manual_stop')
        with self.assertRaisesRegex(ValueError,'stopped'): call(ledger.resolve,self.db,token=reservation['token'],result='sent',min_gap=80,max_gap=80)
        opened=Mock()
        report=stop_controller.finalize(self.db,self.run,self.root,interrupt=False,opener=opened)
        self.assertTrue(report.is_file());self.assertEqual(opened.call_count,1)
        self.assertEqual(stop_controller.finalize(self.db,self.run,self.root,interrupt=False,opener=opened),report)
        self.assertEqual(opened.call_count,1)
        self.assertEqual(self.db.execute('SELECT status FROM candidates').fetchone()[0],'unknown')
        self.assertEqual(ledger.unresolved(self.db),1)
        with self.assertRaisesRegex(ValueError,'stopped'):call(candidate_log.sync,self.db,run=self.run,candidate_id=self.candidate,output_root=self.root)
        wb=load_workbook(report)
        try:
            headers=[c.value for c in wb['候选人简报'][1]]
            self.assertIn('触发方式',headers);self.assertIn('运行结果',wb.sheetnames)
            self.assertEqual(wb['候选人简报'].cell(2,len(headers)).value,'查看名片')
        finally:wb.close()

    def test_zero_candidate_expiry_creates_excel_despite_pause(self):
        call(ledger.control,self.db,run=self.run,action='pause')
        self.db.execute('UPDATE runs SET deadline_at=? WHERE run_id=?',(ledger.iso(ledger.now()-timedelta(seconds=1)),self.run))
        report=stop_controller.poll(self.db,self.run,self.root,interrupt=False,opener=Mock())
        self.assertTrue(report.is_file())
        self.assertEqual(ledger.run_row(self.db,self.run)['stop_reason'],'run_expired')

    def test_stopped_dispatch_cannot_launch_process(self):
        run_safety.request_stop(self.db,self.run,'manual_stop')
        with patch.object(safe_action.subprocess,'Popen') as process:
            with self.assertRaisesRegex(ValueError,'stopped'):safe_action.dispatch(self.db,self.run,0,['never-run'])
            process.assert_not_called()

    def test_missing_chat_binding_blocks_new_policy_before_run_creation(self):
        self.profile.write_text(PROFILE.replace('position: Sales','position: Sales\nevaluation_mode: card\ntrigger_policy: worker'),encoding='utf-8')
        before=self.db.execute('SELECT COUNT(*) FROM runs').fetchone()[0]
        result=call(ledger.start,self.db,profile=str(self.profile),target=150,duration_min=240,min_score=70)
        self.assertEqual(result['reason'],'stop_binding_required')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM runs').fetchone()[0],before)

    def test_stop_serializes_with_admitted_action_and_then_fences_next(self):
        admitted=threading.Event();events=[];errors=[]
        def action():
            try:
                with contextlib.closing(ledger.connect(self.root/'ledger.db')) as db:
                    with run_safety.action_lock(db):
                        run_safety.ensure_writable(db,self.run,0)
                        events.append('admitted');admitted.set();time.sleep(.15);events.append('finished')
            except Exception as exc: errors.append(exc)
        worker=threading.Thread(target=action);worker.start();self.assertTrue(admitted.wait(2))
        run_safety.request_stop(self.db,self.run,'manual_stop');events.append('stopped');worker.join(2)
        self.assertFalse(errors);self.assertEqual(events,['admitted','finished','stopped'])
        with self.assertRaisesRegex(ValueError,'stopped'):run_safety.ensure_writable(self.db,self.run,0)

    def test_managed_helper_cannot_launch_after_stop(self):
        run_safety.request_stop(self.db,self.run,'manual_stop')
        with patch.object(managed_helper.subprocess,'Popen') as process:
            with self.assertRaisesRegex(ValueError,'stopped'):managed_helper.run(self.db,self.run,0,['never-run'])
            process.assert_not_called()

    def test_running_local_helper_is_cancelled_and_deregistered(self):
        errors=[]
        def helper():
            try:
                with contextlib.closing(ledger.connect(self.root/'ledger.db')) as db:
                    managed_helper.run(db,self.run,0,[sys.executable,'-c','import time;time.sleep(5)'])
            except ValueError as exc: errors.append(str(exc))
        worker=threading.Thread(target=helper);worker.start()
        deadline=time.monotonic()+3
        while not self.db.execute('SELECT 1 FROM run_helpers').fetchone():
            if time.monotonic()>deadline:self.fail('helper not registered')
            time.sleep(.03)
        run_safety.request_stop(self.db,self.run,'manual_stop');worker.join(3)
        self.assertFalse(worker.is_alive());self.assertEqual(errors,['helper_cancelled'])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM run_helpers').fetchone()[0],0)

    def test_locked_terminal_filename_uses_separate_snapshot(self):
        run_safety.request_stop(self.db,self.run,'manual_stop')
        replace=candidate_log.os.replace;attempts=[]
        def locked_once(source,target):
            attempts.append(str(target))
            if len(attempts)==1: raise PermissionError('test workbook locked')
            return replace(source,target)
        with patch.object(candidate_log.os,'replace',side_effect=locked_once):
            report=stop_controller.finalize(self.db,self.run,self.root,interrupt=False,opener=Mock())
        self.assertTrue(report.is_file());self.assertIn('独立快照',report.name)

if __name__=='__main__': unittest.main()
