"""Independent, idempotent run termination; never relies on the interrupted agent."""
from __future__ import annotations
import argparse
import contextlib
import io
import json
import os
import sqlite3
import subprocess
import shutil
import sys
import tempfile
import time
import uuid
from pathlib import Path
import ledger
import run_safety
import powershell_host
from windows_process import hidden_options

def load_binding(path):
    binding=json.loads(path.read_text(encoding='utf-8-sig'))
    if binding.get('kind') not in ('uia','app_server'): raise ValueError('invalid_stop_binding_kind')
    return binding

def binding_action(binding,stop=False):
    if binding['kind']=='uia':
        with tempfile.TemporaryDirectory(prefix='codex-stop-') as directory:
            source=Path(directory)/'binding.json';source.write_text(json.dumps(binding,ensure_ascii=False),encoding='utf-8')
            result=subprocess.run(powershell_host.command(Path(__file__).with_name('codex-stop.ps1'),'-Binding',source,'-Action','Stop' if stop else 'Probe'),capture_output=True,timeout=8,**hidden_options())
            if result.returncode: raise RuntimeError('codex_stop_failed: '+result.stderr.decode('utf-8',errors='replace')[-1600:])
            return json.loads(result.stdout.decode('utf-8-sig'))
    with tempfile.TemporaryDirectory(prefix='codex-stop-') as directory:
        source=Path(directory)/'binding.json'
        source.write_text(json.dumps(binding,ensure_ascii=False),encoding='utf-8')
        result=subprocess.run(powershell_host.command(Path(__file__).with_name('app-server-stop.ps1'),'-Binding',source,'-Action','Stop' if stop else 'Probe'),capture_output=True,timeout=6,**hidden_options())
        if result.returncode: raise RuntimeError('app_server_stop_failed: '+result.stderr.decode('utf-8',errors='replace')[-1200:])
        return json.loads(result.stdout.decode('utf-8-sig'))


def probe_binding(binding): return binding_action(binding)


def bind_run(db,run_id,chat_title=None):
    run=ledger.run_row(db,run_id)
    db.execute('CREATE TABLE IF NOT EXISTS overlay_chat(run_id TEXT PRIMARY KEY,chat_title TEXT NOT NULL)')
    saved=db.execute('SELECT chat_title FROM overlay_chat WHERE run_id=?',(run_id,)).fetchone()
    if chat_title:
        db.execute('INSERT OR REPLACE INTO overlay_chat VALUES (?,?)',(run_id,chat_title))
    elif saved:chat_title=saved[0]
    existing=json.loads(run['stop_binding']) if run['stop_binding'] else None
    if existing and not chat_title:
        if existing['kind']=='app_server' or not existing.get('chat_title'):return existing
        chat_title=existing['chat_title']
    if not chat_title:
        _,_,fields=ledger.profile_check(Path(run['profile_path']))
        date=ledger.parse_time(run['started_at']).astimezone(ledger.timezone(ledger.timedelta(hours=8))).date()
        chat_title=f"{fields['company']} {fields['position']} - {date.isoformat()}"
    db.execute('INSERT OR REPLACE INTO overlay_chat VALUES (?,?)',(run_id,chat_title))
    with tempfile.TemporaryDirectory(prefix='codex-bind-') as directory:
        path=Path(directory)/'binding.json'
        command=powershell_host.command(Path(__file__).with_name('bind-codex-stop.ps1'),'-ChatMarker',chat_title,'-Output',path)
        result=subprocess.run(command,capture_output=True,timeout=8,**hidden_options())
        if result.returncode:raise RuntimeError('codex_binding_failed: '+result.stderr.decode('utf-8',errors='replace')[-1200:])
        binding=load_binding(path)
    db.execute('UPDATE runs SET stop_binding=? WHERE run_id=?',(json.dumps(binding,ensure_ascii=False),run_id))
    return binding


def interrupt_run(db,run_id):
    try:
        binding=bind_run(db,run_id)
        status=binding_action(binding,True)['state']
        if status not in ('interrupted','idle'):raise RuntimeError('unexpected_interrupt_state: '+status)
    except Exception as exc:
        status='failed: '+str(exc)
    db.execute('UPDATE runs SET interrupt_status=? WHERE run_id=?',(status,run_id))
    return status

def final_report(db,run_id,root):
    import candidate_log
    run=ledger.run_row(db,run_id)
    # Read from an isolated SQLite snapshot so a later manual reconciliation cannot
    # change the terminal workbook while it is being rendered.
    snapshot=sqlite3.connect(':memory:');snapshot.row_factory=sqlite3.Row
    db.backup(snapshot)
    try:
        name=f'{Path(candidate_log.report_filename(run["profile_path"])).stem}-结束-{run_id[:8]}-{uuid.uuid4().hex[:6]}.xlsx'
        try:
            report=candidate_log.workbook(snapshot,root,run['started_at'],filename=name,profile_path=run['profile_path'])
        except PermissionError:
            alternate=Path(name).stem+'-独立快照-'+uuid.uuid4().hex[:8]+'.xlsx'
            report=candidate_log.workbook(snapshot,root,run['started_at'],filename=alternate,profile_path=run['profile_path'])
        from openpyxl import load_workbook
        wb=load_workbook(report)
        summary=wb.create_sheet('运行结果')
        summary.append(['运行编号',run_id]);summary.append(['停止原因',run['stop_reason']])
        summary.append(['Codex中断结果',run['interrupt_status'] or '未核验'])
        summary.append(['运行时间限制','无'])
        summary.append(['成功招呼人数',snapshot.execute("SELECT COUNT(*) FROM candidates WHERE run_id=? AND status='sent'",(run_id,)).fetchone()[0]])
        summary.append(['查看人数',snapshot.execute('SELECT COUNT(*) FROM candidate_progress WHERE run_id=?',(run_id,)).fetchone()[0]])
        summary.append(['未完成候选人（哈希）','状态'])
        for row in snapshot.execute('SELECT candidate_hash FROM candidate_progress WHERE run_id=? AND finished_at IS NULL',(run_id,)):
            summary.append([row[0],'interrupted；未完成，不伪造判定或离开时间'])
        summary.column_dimensions['A'].width=42;summary.column_dimensions['B'].width=86
        try: wb.save(report)
        except PermissionError:
            report=report.with_name(report.stem+'-'+uuid.uuid4().hex[:8]+'.xlsx');wb.save(report)
        wb.close()
        return report
    finally: snapshot.close()

def finalize(db,run_id,root,interrupt=True,opener=None):
    import candidate_log
    candidate_log.table(db)
    with run_safety.action_lock(db):
        run=ledger.run_row(db,run_id)
        if run['finalized_at']: return Path(run['report_path']) if run['report_path'] else None
        if not run['stop_reason']: raise ValueError('stop_must_precede_finalize')
        if run['finalizer_owner']:
            parts=run['finalizer_owner'].split(':')
            if len(parts)!=3: raise ValueError('unrecognized_finalizer_owner_requires_manual_recovery')
            try:
                if abs(run_safety.process_birth(int(parts[0]))-float(parts[1]))<.0001: return None
            except (ProcessLookupError,OSError): pass
        owner=f'{os.getpid()}:{run_safety.process_birth(os.getpid()):.9f}:{uuid.uuid4().hex}'
        db.execute('UPDATE runs SET finalizer_owner=?,finalizer_claimed_at=? WHERE run_id=?',(owner,ledger.iso(ledger.now()),run_id))
    try:
        if interrupt:interrupt_run(db,run_id)
        # Cancel only registered helper processes whose creation time still matches.
        until=time.monotonic()+5
        for row in db.execute('SELECT pid,created_at FROM run_helpers WHERE run_id=?',(run_id,)):
            try:
                if time.monotonic()>=until: break
                run_safety.process_birth(row['pid'],terminate=True,expected=row['created_at'])
            except (ProcessLookupError,OSError,ValueError): pass
        with run_safety.action_lock(db):
            at=ledger.iso(ledger.now())
            db.execute("UPDATE message_attempts SET status='unknown',updated_at=? WHERE run_id=? AND status IN ('pending','dispatching')",(at,run_id))
            db.execute("UPDATE candidates SET status='unknown',updated_at=? WHERE run_id=? AND status='pending'",(at,run_id))
            db.execute("UPDATE assessments SET send_status=COALESCE((SELECT status FROM candidates c WHERE c.candidate_hash=assessments.candidate_hash),'interrupted') WHERE run_id=?",(run_id,))
            report=final_report(db,run_id,root)
            db.execute('UPDATE runs SET finalized_at=?,report_path=? WHERE run_id=?',(at,str(report),run_id))
        if opener is not None: opener(str(report))
        elif os.name=='nt': os.startfile(report)
        return report
    except Exception:
        db.execute('UPDATE runs SET finalizer_owner=NULL WHERE run_id=? AND finalized_at IS NULL AND finalizer_owner=?',(run_id,owner))
        raise

def poll(db,run_id,root,interrupt=True,opener=None):
    run=ledger.run_row(db,run_id)
    reason=run['stop_reason']
    if not reason and run_safety.quantity_reason(db,run_id) and ledger.active_progress(db,run_id) is None:
        reason=run_safety.quantity_reason(db,run_id)
    if reason:
        run_safety.request_stop(db,run_id,reason)
        return finalize(db,run_id,root,interrupt,opener)
    return None

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--db',type=Path,default=ledger.DEFAULT_DB)
    p.add_argument('--run',required=True);p.add_argument('--output-root',type=Path,default=Path.home()/'Desktop')
    p.add_argument('--stop',action='store_true');a=p.parse_args()
    with contextlib.closing(ledger.connect(a.db)) as db:
        if a.stop: run_safety.request_stop(db,a.run,'manual_stop')
        while not poll(db,a.run,a.output_root): time.sleep(1)
if __name__=='__main__': main()
