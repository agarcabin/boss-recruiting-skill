"""Persist each candidate decision and rebuild the daily recruiting workbook."""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

import ledger
import coworker_progress
import recruiting_rules
import run_safety
from report_paths import profile_identity, report_filename


try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass

CN_TZ = timezone(timedelta(hours=8))
SUMMARY_HEADERS = ["录入时间", "候选人", "综合分", "判定", "判断时长", "优点", "劣势", "触发方式", "招呼结果", "查看简历"]
DETAIL_HEADERS = ["录入时间", "项目", "岗位", "运行编号", "候选人", "候选人标识哈希", "开始时间", "结束时间",
                  "判断时长", "判断时长（秒）", "列表分", "列表依据", "电子简历分", "电子简历依据",
                  "附件简历分", "附件依据", "综合分", "硬条件", "硬条件例外授权", "判定", "发送状态", "优点", "劣势",
                  "电子简历判断", "PDF归档", "PDF路径", "PDF说明",
                  "简历图片", "简历OCR文本", "简历OCR结果", "简历归档", "简历归档说明",
                  "同事去重", "首条同事招呼时间", "同事去重依据"]
GREEN_FILL = "C6EFCE"
GREEN_FONT = "006100"
YELLOW_FILL = "FFEB9C"
YELLOW_FONT = "9C6500"


def output(**data):
    print(json.dumps(data, ensure_ascii=False))


def table(db):
    db.execute("""CREATE TABLE IF NOT EXISTS assessments (
        run_id TEXT NOT NULL, candidate_hash TEXT NOT NULL, candidate_name TEXT NOT NULL,
        started_at TEXT NOT NULL, ended_at TEXT NOT NULL, list_score INTEGER,
        list_evidence TEXT, electronic_score INTEGER, electronic_evidence TEXT,
        attachment_score INTEGER, attachment_evidence TEXT, overall_score INTEGER NOT NULL,
        hard_requirements TEXT NOT NULL, decision TEXT NOT NULL, pros TEXT NOT NULL,
        cons TEXT NOT NULL, electronic_enabled INTEGER NOT NULL, send_status TEXT NOT NULL,
        pdf_status TEXT NOT NULL DEFAULT 'not_applicable', pdf_path TEXT,
        pdf_note TEXT, updated_at TEXT NOT NULL, report_written INTEGER NOT NULL DEFAULT 0,
        hard_requirement_override TEXT NOT NULL DEFAULT 'false',
        hard_requirement_override_note TEXT,
        PRIMARY KEY(run_id,candidate_hash))""")
    columns = {
        row["name"]
        for row in db.execute("PRAGMA table_info(assessments)").fetchall()
    }
    additions = {
        "resume_image_path": "TEXT",
        "resume_ocr_text_path": "TEXT",
        "resume_ocr_json_path": "TEXT",
        "resume_archive_status": "TEXT NOT NULL DEFAULT 'not_applicable'",
        "resume_archive_note": "TEXT",
        "assessed_at": "TEXT",
        "hard_requirement_override": "TEXT NOT NULL DEFAULT 'false'",
        "hard_requirement_override_note": "TEXT",
        "coworker_progress": "TEXT",
        "evaluation_mode": "TEXT NOT NULL DEFAULT 'resume'",
        "eligibility_json": "TEXT",
        "evaluation_input": "TEXT",
        "card_image_path": "TEXT",
        "card_uia_path": "TEXT",
        "timing_source": "TEXT",
    }
    for name, definition in additions.items():
        if name not in columns:
            db.execute(f"ALTER TABLE assessments ADD COLUMN {name} {definition}")
    db.execute("""UPDATE assessments SET resume_archive_status='pending'
        WHERE send_status='sent' AND resume_archive_status='not_applicable' AND evaluation_mode<>'card'""")


def daily_dir(root: Path, started_at: str) -> Path:
    stamp = datetime.fromisoformat(started_at).astimezone(CN_TZ)
    return root / "自动投递记录" / f"{stamp.month}月{stamp.day}日"


def recorded_time(assessed_at: str) -> str:
    return datetime.fromisoformat(assessed_at).astimezone(CN_TZ).strftime("%H:%M")


def judgment_seconds(started_at: str, ended_at: str) -> int | None:
    if not ended_at:
        return None
    start = datetime.fromisoformat(started_at)
    end = datetime.fromisoformat(ended_at)
    return max(0, int((end - start).total_seconds()))


def judgment_duration_text(started_at: str, ended_at: str) -> str:
    seconds = judgment_seconds(started_at, ended_at)
    if seconds is None:
        return "待完成"
    return f"{seconds // 60}分{seconds % 60}秒"


def decision_label(decision: str) -> str:
    return {"greet": "打招呼", "review": "备选复核", "skip": "不匹配"}.get(decision, decision)


def decision_style(decision: str) -> tuple[str, str] | None:
    if decision == "greet":
        return GREEN_FILL, GREEN_FONT
    if decision == "review":
        return YELLOW_FILL, YELLOW_FONT
    return None


def score_source(payload: dict, key: str, enabled: bool) -> tuple[int | None, str]:
    source = payload.get(key)
    if source is None:
        return None, "未查看" if enabled else "已关闭判断"
    if not isinstance(source, dict) or not isinstance(source.get("score"), int) or not 0 <= source["score"] <= 100:
        raise ValueError(f"{key} needs integer score 0..100")
    evidence = str(source.get("evidence", "")).strip()
    if not evidence:
        raise ValueError(f"{key} evidence required")
    return source["score"], evidence


def validate(payload: dict) -> dict:
    for key in ("candidate_id", "candidate_name", "started_at", "pros", "cons", "decision", "overall_score", "hard_requirements", "electronic_resume_enabled"):
        if key not in payload:
            raise ValueError(f"missing {key}")
    if payload["decision"] not in ("greet", "skip", "review"):
        raise ValueError("decision must be greet/skip/review")
    if type(payload["overall_score"]) is not int or not 0 <= payload["overall_score"] <= 100:
        raise ValueError("overall_score must be integer 0..100")
    if type(payload["electronic_resume_enabled"]) is not bool:
        raise ValueError("electronic_resume_enabled must be boolean")
    if type(payload["hard_requirements"]) is not bool:
        raise ValueError("hard_requirements must be boolean")
    hard_override = payload.get("hard_requirement_override", False)
    if type(hard_override) is not bool:
        raise ValueError("hard_requirement_override must be boolean")
    hard_override_note = str(payload.get("hard_requirement_override_note", "")).strip()
    if hard_override and not hard_override_note:
        raise ValueError("hard_requirement_override_note required when override is enabled")
    start = datetime.fromisoformat(payload["started_at"])
    if start.tzinfo is None:
        raise ValueError("started_at must include timezone")
    if payload.get("ended_at") not in (None, ""):
        raise ValueError("ended_at is set after leaving the candidate; use finalize-timing")
    if not str(payload["candidate_name"]).strip() or not str(payload["pros"]).strip() or not str(payload["cons"]).strip():
        raise ValueError("candidate_name, pros and cons required; use '未见明显...' when justified")
    list_score, list_evidence = score_source(payload, "list", False)
    if payload["electronic_resume_enabled"] and not payload.get("electronic"):
        raise ValueError("electronic assessment required when electronic_resume_enabled is true")
    electronic_score, electronic_evidence = score_source(payload, "electronic", payload["electronic_resume_enabled"])
    if not payload["electronic_resume_enabled"] and electronic_score is not None:
        raise ValueError("electronic score cannot be supplied when judgement disabled")
    if payload["electronic_resume_enabled"] and payload["overall_score"] != electronic_score:
        raise ValueError("overall_score must equal electronic score")
    attachment_score, attachment_evidence = score_source(payload, "attachment", True)
    if payload["decision"] == "greet" and not payload["hard_requirements"] and not hard_override and payload.get("trigger_policy","legacy") not in ("worker","agv_advanced","agv_basic"):
        raise ValueError("greet requires hard requirements met")
    if hard_override and payload["decision"] != "greet":
        raise ValueError("hard requirement override is only valid for an authorized greeting")
    evidence = payload.get("resume_evidence")
    if payload["electronic_resume_enabled"] and not evidence:
        raise ValueError("resume_evidence required when electronic_resume_enabled is true")
    if evidence is not None:
        if not isinstance(evidence, dict):
            raise ValueError("resume_evidence must be an object")
        if not str(evidence.get("image", "")).strip() or not str(evidence.get("ocr_text", "")).strip():
            raise ValueError("resume_evidence image and ocr_text required")
    coworker = ({"may_continue": True, "action": "not_applicable", "evidence": "列表模式不检查详情同事记录"}
                if payload.get("evaluation_mode")=="card" else coworker_progress.evaluate(payload.get("coworker_progress"), ledger.now()))
    if payload["decision"] == "greet" and not coworker["may_continue"]:
        raise ValueError(coworker["reason"])
    return dict(list_score=list_score, list_evidence=list_evidence,
                electronic_score=electronic_score, electronic_evidence=electronic_evidence,
                attachment_score=attachment_score, attachment_evidence=attachment_evidence,
                resume_evidence=evidence, hard_requirement_override=hard_override,
                hard_requirement_override_note=hard_override_note, coworker_progress=coworker)


def workbook(db, root: Path, started_at: str, filename: str | None = None,
             profile_path: Path | str | None = None) -> Path:
    folder = daily_dir(root, started_at)
    folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "候选人简报"
    headers=list(SUMMARY_HEADERS)
    if profile_path and ledger.profile_check(Path(profile_path))[2].get("evaluation_mode")=="card": headers[-1]="查看名片"
    summary.append(headers)
    details = wb.create_sheet("详细信息")
    details.append(DETAIL_HEADERS)
    details.sheet_state = "hidden"
    rows = db.execute("""SELECT a.*, r.profile_path FROM assessments a
        JOIN runs r ON r.run_id=a.run_id WHERE a.started_at>=? AND a.started_at<?
        ORDER BY a.started_at,a.candidate_name""", (
            datetime.fromisoformat(started_at).astimezone(CN_TZ).replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc).isoformat(),
            (datetime.fromisoformat(started_at).astimezone(CN_TZ).replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1)).astimezone(timezone.utc).isoformat(),
        )).fetchall()
    if profile_path is None:
        identities = {profile_identity(row["profile_path"]) for row in rows}
        if len(identities) != 1:
            raise ValueError("report_profile_required_for_multiple_or_no_projects")
        profile_path = rows[0]["profile_path"]
    identity = profile_identity(profile_path)
    path = folder / report_filename(profile_path, filename)
    for row in rows:
        if profile_identity(row["profile_path"]) != identity:
            continue
        _, _, fields = ledger.profile_check(Path(row["profile_path"]))
        entered_at = recorded_time(row["assessed_at"] or row["ended_at"] or row["updated_at"])
        duration = judgment_duration_text(row["started_at"], row["ended_at"])
        duration_seconds = judgment_seconds(row["started_at"], row["ended_at"])
        coworker = json.loads(row["coworker_progress"] or "{}")
        coworker_skipped = coworker.get("action") == "skip" and row["decision"] == "skip"
        label = "重复跳过" if coworker_skipped else decision_label(row["decision"])
        overall_score = None if coworker_skipped else row["overall_score"]
        summary.append([entered_at, row["candidate_name"], overall_score, label, duration,
                        row["pros"], row["cons"], json.loads(row["eligibility_json"] or "{}").get("trigger","score"),
                        message_summary(db,row), "查看名片" if row["evaluation_mode"]=="card" else "查看简历"])
        link_cell = summary.cell(row=summary.max_row, column=len(SUMMARY_HEADERS))
        evidence_image=row["card_image_path"] if row["evaluation_mode"]=="card" else row["resume_image_path"]
        if evidence_image and Path(evidence_image).is_file():
            link_cell.hyperlink = Path(evidence_image).resolve().relative_to(folder.resolve()).as_posix()
            link_cell.style = "Hyperlink"
        else:
            link_cell.value = "未保存"
        details.append([entered_at, fields.get("company", ""), fields.get("position", ""), row["run_id"],
                        row["candidate_name"], row["candidate_hash"], row["started_at"], row["ended_at"],
                        duration, duration_seconds, row["list_score"], row["list_evidence"],
                        row["electronic_score"], row["electronic_evidence"], row["attachment_score"],
                        row["attachment_evidence"], overall_score,
                        "符合" if row["hard_requirements"] == "true" else "未确认/不符",
                        row["hard_requirement_override_note"] if row["hard_requirement_override"] == "true" else "否",
                        label,
                        row["send_status"], row["pros"], row["cons"],
                        "未评估（同事去重）" if coworker_skipped else ("开启" if row["electronic_enabled"] else "关闭"),
                        row["pdf_status"], row["pdf_path"] or "", row["pdf_note"] or "",
                        row["resume_image_path"] or "", row["resume_ocr_text_path"] or "",
                        row["resume_ocr_json_path"] or "", row["resume_archive_status"],
                        row["resume_archive_note"] or "", coworker.get("action", "未核验"),
                        coworker.get("first_greeting_at") or "", coworker.get("evidence", "")])
        style = decision_style(row["decision"])
        if style:
            fill, font_color = style
            for cell in summary[summary.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(color=font_color)
            for cell in details[details.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(color=font_color)
    evidence_sheet=wb.create_sheet("评估证据")
    evidence_sheet.sheet_state="hidden"
    evidence_sheet.append(["运行编号","候选人哈希","模式","触发依据","三句状态","卡片图","UIA原文","计时来源"])
    for item in rows:
        if profile_identity(item["profile_path"])==identity:
            attempts=[dict(x) for x in db.execute("SELECT message_index,status FROM message_attempts WHERE run_id=? AND candidate_hash=? ORDER BY message_index",(item["run_id"],item["candidate_hash"]))]
            evidence_sheet.append([item["run_id"],item["candidate_hash"],item["evaluation_mode"],item["eligibility_json"],json.dumps(attempts,ensure_ascii=False),item["card_image_path"],item["card_uia_path"],item["timing_source"]])
    for ws in (summary, details):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="26496A")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    summary.column_dimensions["A"].width = 12
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 12
    summary.column_dimensions["D"].width = 13
    summary.column_dimensions["E"].width = 14
    summary.column_dimensions["F"].width = 48
    summary.column_dimensions["G"].width = 48
    summary.column_dimensions["H"].width = 14
    for hidden_header in ("触发方式", "招呼结果"):
        summary.column_dimensions[
            get_column_letter(headers.index(hidden_header) + 1)
        ].hidden = True
    for col in details.columns:
        letter = col[0].column_letter
        details.column_dimensions[letter].width = min(
            52, max(13, max(len(str(c.value or "")) for c in col[:100]) + 2)
        )
    handle, tmp = tempfile.mkstemp(prefix=".简报-", suffix=".xlsx", dir=folder)
    os.close(handle)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return path


def record(db, args, payload=None):
    if payload is None:
        payload = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
    run = ledger.run_row(db,args.run)
    rules = recruiting_rules.run_rules(run)
    payload = dict(payload)
    payload["evaluation_mode"] = rules.get("mode","resume")
    payload["trigger_policy"] = rules.get("policy","legacy")
    result = recruiting_rules.evaluate(rules,payload,run["min_score"])
    if rules.get("policy","legacy") != "legacy":
        payload["overall_score"] = result["score"]
        if payload.get("electronic"): payload["electronic"] = dict(payload["electronic"],score=result["score"])
        if rules.get("mode")=="card":
            if payload["electronic_resume_enabled"]: raise ValueError("card_must_not_open_resume")
            payload["list"] = {"score":result["score"],"evidence":json.dumps(result["hits"],ensure_ascii=False) or "未命中"}
        elif not payload["electronic_resume_enabled"]: raise ValueError("resume_policy_requires_electronic")
    if payload["decision"]=="greet" and not result["eligible"]:
        raise ValueError(f"greet requires score >= {run['min_score']} or valid keyword and permitted requirements")
    scores = validate(payload)
    payload["started_at"] = datetime.fromisoformat(payload["started_at"]).astimezone(timezone.utc).isoformat(timespec="seconds")
    run = ledger.run_row(db, args.run)
    if payload["decision"] == "greet" and not result["eligible"]:
        raise ValueError("assessment_not_eligible")
    _, _, fields = ledger.profile_check(Path(run["profile_path"]))
    digest = ledger.run_candidate_hash(db,args.run,payload["candidate_id"])
    active = ledger.active_progress(db, args.run)
    if not active or active["candidate_hash"] != digest:
        raise ValueError("candidate_not_active; run ledger.py candidate-begin first")
    if db.execute("SELECT 1 FROM assessments WHERE run_id=? AND candidate_hash=?", (args.run, digest)).fetchone():
        raise ValueError("candidate already assessed; use sync for status changes")
    resume_paths = [None, None, None]
    if scores["resume_evidence"]:
        resume_paths = archive_resume_files(
            args.output_root,
            payload["started_at"],
            str(payload["candidate_name"]).strip(),
            digest,
            scores["resume_evidence"],
        )
    card_paths = archive_card_files(args.output_root,payload["started_at"],payload["candidate_name"],digest,payload) if rules.get("mode")=="card" else (None,None)
    if scores["resume_evidence"]:
        payload["resume_evidence"] = dict(image=resume_paths[0],ocr_text=resume_paths[1],ocr_json=resume_paths[2])
    status = "assessed" if payload["decision"] == "greet" else payload["decision"]
    stamp = ledger.iso(ledger.now())
    db.execute("""INSERT INTO assessments (
        run_id,candidate_hash,candidate_name,started_at,ended_at,list_score,list_evidence,
        electronic_score,electronic_evidence,attachment_score,attachment_evidence,overall_score,
        hard_requirements,hard_requirement_override,hard_requirement_override_note,
        decision,pros,cons,electronic_enabled,send_status,pdf_status,pdf_path,
        pdf_note,updated_at,report_written,resume_image_path,resume_ocr_text_path,
        resume_ocr_json_path,resume_archive_status,assessed_at
    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (args.run, digest, str(payload["candidate_name"]).strip(), payload["started_at"], "",
         scores["list_score"], scores["list_evidence"], scores["electronic_score"], scores["electronic_evidence"],
         scores["attachment_score"], scores["attachment_evidence"], payload["overall_score"],
         json.dumps(payload["hard_requirements"]), json.dumps(scores["hard_requirement_override"]),
         scores["hard_requirement_override_note"] or None, payload["decision"], payload["pros"], payload["cons"],
         int(payload["electronic_resume_enabled"]), status,
         "pending" if status == "assessed" else "not_applicable", None, None, stamp, 0,
         resume_paths[0], resume_paths[1], resume_paths[2],
         "saved" if resume_paths[0] else "not_applicable", stamp))
    db.execute("UPDATE assessments SET evaluation_mode=?,eligibility_json=?,evaluation_input=?,card_image_path=?,card_uia_path=?,timing_source=? WHERE run_id=? AND candidate_hash=?",
               (rules.get("mode","resume"),json.dumps(result,ensure_ascii=False),json.dumps(payload,ensure_ascii=False),*card_paths,
                "card_read" if rules.get("mode")=="card" else "detail_click",args.run,digest))
    if rules.get("mode")=="card" or not rules.get("pdf_required",True):
        db.execute("UPDATE assessments SET pdf_status='not_applicable' WHERE run_id=? AND candidate_hash=?",(args.run,digest))
    db.execute("UPDATE assessments SET coworker_progress=? WHERE run_id=? AND candidate_hash=?",
               (json.dumps(scores["coworker_progress"], ensure_ascii=False), args.run, digest))
    path = workbook(db, args.output_root, payload["started_at"], profile_path=run["profile_path"])
    db.execute("UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?", (args.run, digest))
    output(state="recorded", candidate_hash=digest[:12], report=str(path), decision=payload["decision"])


def finalize_timing(db, args):
    digest = ledger.run_candidate_hash(db, args.run, args.candidate_id)
    ended_at = ledger.parse_time(args.ended_at)
    at = ledger.now()
    if ended_at > at + timedelta(seconds=5):
        raise ValueError("candidate_end_in_future")
    db.execute("BEGIN IMMEDIATE")
    try:
        active = ledger.active_progress(db, args.run)
        if not active or active["candidate_hash"] != digest:
            raise ValueError("candidate_not_active")
        row = db.execute(
            "SELECT started_at,ended_at FROM assessments WHERE run_id=? AND candidate_hash=?",
            (args.run, digest),
        ).fetchone()
        if not row:
            raise ValueError("assessment_missing")
        if ended_at < ledger.parse_time(active["started_at"]):
            raise ValueError("candidate_end_before_start")
        if row["ended_at"]:
            if ledger.iso(ledger.parse_time(row["ended_at"])) != ledger.iso(ended_at):
                raise ValueError("candidate_timing_already_finalized")
            db.commit()
            output(state="timing_already_finalized", candidate_ended_at=ledger.iso(ended_at))
            return
        stamp = ledger.iso(ended_at)
        db.execute(
            "UPDATE assessments SET ended_at=?,updated_at=?,report_written=0 "
            "WHERE run_id=? AND candidate_hash=?",
            (stamp, ledger.iso(at), args.run, digest),
        )
        path = workbook(db, args.output_root, row["started_at"],
                        profile_path=ledger.run_row(db, args.run)["profile_path"])
        db.execute(
            "UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?",
            (args.run, digest),
        )
        db.commit()
    except Exception:
        db.rollback()
        raise
    output(state="timing_finalized", candidate_ended_at=stamp, report=str(path))


def sync(db, args):
    run = ledger.run_row(db, args.run)
    _, _, fields = ledger.profile_check(Path(run["profile_path"]))
    digest = ledger.run_candidate_hash(db,args.run,args.candidate_id)
    row = db.execute("SELECT * FROM assessments WHERE run_id=? AND candidate_hash=?", (args.run, digest)).fetchone()
    if not row:
        raise ValueError("assessment_missing")
    candidate = db.execute("SELECT status FROM candidates WHERE candidate_hash=?", (digest,)).fetchone()
    if not candidate:
        raise ValueError("ledger_status_missing")
    status = candidate["status"]
    pdf_status = "pending" if status == "sent" else ("not_applicable" if status in ("skipped", "review", "not-sent") else row["pdf_status"])
    if row["evaluation_mode"]=="card" or not recruiting_rules.run_rules(run).get("pdf_required",True): pdf_status="not_applicable"
    elif row["pdf_status"] in ("saved","missing"): pdf_status=row["pdf_status"]
    resume_status = row["resume_archive_status"]
    if resume_status not in ("saved", "missing"):
        resume_status = "pending" if status == "sent" else "not_applicable"
    if row["evaluation_mode"]=="card": resume_status="not_applicable"
    db.execute("""UPDATE assessments SET send_status=?,pdf_status=?,resume_archive_status=?,
        updated_at=?,report_written=0 WHERE run_id=? AND candidate_hash=?""",
               (status, pdf_status, resume_status, ledger.iso(ledger.now()), args.run, digest))
    path = workbook(db, args.output_root, row["started_at"], profile_path=run["profile_path"])
    db.execute("UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?", (args.run, digest))
    output(state="synced", send_status=status, report=str(path))


def archive(db, args):
    run = ledger.run_row(db, args.run)
    _, _, fields = ledger.profile_check(Path(run["profile_path"]))
    digest = ledger.candidate_hash(f"{fields['company']}-{fields['position']}", args.candidate_id)
    row = db.execute("SELECT * FROM assessments WHERE run_id=? AND candidate_hash=?", (args.run, digest)).fetchone()
    if not row or row["send_status"] != "sent":
        raise ValueError("only confirmed sent candidate can be archived")
    source = Path(args.pdf)
    if not source.is_file() or source.stat().st_size < 100:
        raise ValueError("pdf_missing_or_invalid")
    with source.open("rb") as stream:
        header = stream.read(5)
    if header != b"%PDF-":
        raise ValueError("pdf_missing_or_invalid")
    try:
        if len(PdfReader(str(source)).pages) < 1:
            raise ValueError("pdf_has_no_pages")
    except Exception as exc:
        raise ValueError(f"pdf_unreadable: {exc}") from exc
    folder = daily_dir(args.output_root, row["started_at"])
    folder.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", row["candidate_name"]).strip(" .")[:40] or "候选人"
    target = folder / f"{safe_name}-{digest[:10]}.pdf"
    if target.exists():
        if hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(source.read_bytes()).digest():
            raise ValueError("archive_filename_collision")
    else:
        handle, tmp = tempfile.mkstemp(prefix=".resume-", suffix=".pdf", dir=folder)
        os.close(handle)
        try:
            shutil.copyfile(source, tmp)
            os.replace(tmp, target)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
    db.execute("UPDATE assessments SET pdf_status='saved',pdf_path=?,pdf_note='',updated_at=?,report_written=0 WHERE run_id=? AND candidate_hash=?",
               (str(target), ledger.iso(ledger.now()), args.run, digest))
    path = workbook(db, args.output_root, row["started_at"], profile_path=run["profile_path"])
    db.execute("UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?", (args.run, digest))
    output(state="archived", pdf=str(target), report=str(path))


def safe_candidate_name(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")[:40] or "候选人"


def copy_atomic(source: Path, target: Path) -> None:
    if target.exists():
        if hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(source.read_bytes()).digest():
            raise ValueError("archive_filename_collision")
        return
    handle, temporary = tempfile.mkstemp(prefix=".resume-", suffix=target.suffix, dir=target.parent)
    os.close(handle)
    try:
        shutil.copyfile(source, temporary)
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def archive_resume_files(
    root: Path,
    started_at: str,
    candidate_name: str,
    digest: str,
    evidence: dict,
) -> tuple[str, str, str | None]:
    image = Path(str(evidence.get("image", "")))
    ocr_text = Path(str(evidence.get("ocr_text", "")))
    ocr_json_value = str(evidence.get("ocr_json", "")).strip()
    ocr_json = Path(ocr_json_value) if ocr_json_value else None

    if not image.is_file() or image.stat().st_size < 100:
        raise ValueError("resume_image_missing_or_invalid")
    with image.open("rb") as stream:
        if stream.read(8) != b"\x89PNG\r\n\x1a\n":
            raise ValueError("resume_image_missing_or_invalid")
    if not ocr_text.is_file() or len(re.sub(r"\s+", "", ocr_text.read_text(encoding="utf-8-sig"))) < 100:
        raise ValueError("resume_ocr_text_missing_or_invalid")
    if ocr_json:
        if not ocr_json.is_file():
            raise ValueError("resume_ocr_json_missing")
        try:
            json.loads(ocr_json.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ValueError("resume_ocr_json_invalid") from exc

    folder = daily_dir(root, started_at)
    image_folder = folder / "简历图片"
    ocr_folder = folder / "简历OCR"
    image_folder.mkdir(parents=True, exist_ok=True)
    ocr_folder.mkdir(parents=True, exist_ok=True)
    safe_name = safe_candidate_name(candidate_name)
    target_image = image_folder / f"{safe_name}-{digest[:10]}.png"
    target_text = ocr_folder / f"{safe_name}-{digest[:10]}.ocr.txt"
    target_json = ocr_folder / f"{safe_name}-{digest[:10]}.ocr.json"
    copy_atomic(image, target_image)
    copy_atomic(ocr_text, target_text)
    if ocr_json:
        copy_atomic(ocr_json, target_json)
    return str(target_image), str(target_text), str(target_json) if ocr_json else None


def archive_resume(db, args):
    run = ledger.run_row(db, args.run)
    _, _, fields = ledger.profile_check(Path(run["profile_path"]))
    digest = ledger.candidate_hash(f"{fields['company']}-{fields['position']}", args.candidate_id)
    row = db.execute("SELECT * FROM assessments WHERE run_id=? AND candidate_hash=?", (args.run, digest)).fetchone()
    if not row:
        raise ValueError("assessment_missing")
    image, text, json_path = archive_resume_files(
        args.output_root,
        row["started_at"],
        row["candidate_name"],
        digest,
        {
            "image": args.image,
            "ocr_text": args.ocr_text,
            "ocr_json": args.ocr_json,
        },
    )
    db.execute("""UPDATE assessments SET resume_image_path=?,resume_ocr_text_path=?,
        resume_ocr_json_path=?,resume_archive_status='saved',resume_archive_note='',
        updated_at=?,report_written=0 WHERE run_id=? AND candidate_hash=?""",
               (image, text, json_path,
                ledger.iso(ledger.now()), args.run, digest))
    path = workbook(db, args.output_root, row["started_at"], profile_path=run["profile_path"])
    db.execute("UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?", (args.run, digest))
    output(state="resume_archived", image=image, ocr_text=text,
           ocr_json=json_path, report=str(path))


def missing_resume(db, args):
    run = ledger.run_row(db, args.run)
    if args.candidate_hash:
        digest = args.candidate_hash.lower()
    else:
        _, _, fields = ledger.profile_check(Path(run["profile_path"]))
        digest = ledger.candidate_hash(f"{fields['company']}-{fields['position']}", args.candidate_id)
    row = db.execute("SELECT * FROM assessments WHERE run_id=? AND candidate_hash=?", (args.run, digest)).fetchone()
    if not row:
        raise ValueError("assessment_missing")
    if not args.reason.strip():
        raise ValueError("missing resume evidence reason required")
    db.execute("""UPDATE assessments SET resume_image_path=NULL,resume_ocr_text_path=NULL,
        resume_ocr_json_path=NULL,resume_archive_status='missing',resume_archive_note=?,
        updated_at=?,report_written=0 WHERE run_id=? AND candidate_hash=?""",
               (args.reason.strip(), ledger.iso(ledger.now()), args.run, digest))
    path = workbook(db, args.output_root, row["started_at"], profile_path=run["profile_path"])
    db.execute("UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?", (args.run, digest))
    output(state="resume_missing_recorded", report=str(path))


def missing(db, args):
    run = ledger.run_row(db, args.run)
    _, _, fields = ledger.profile_check(Path(run["profile_path"]))
    digest = ledger.candidate_hash(f"{fields['company']}-{fields['position']}", args.candidate_id)
    row = db.execute("SELECT * FROM assessments WHERE run_id=? AND candidate_hash=?", (args.run, digest)).fetchone()
    if not row or row["send_status"] != "sent":
        raise ValueError("only confirmed sent candidate can have missing PDF")
    if not args.reason.strip():
        raise ValueError("missing PDF reason required")
    db.execute("UPDATE assessments SET pdf_status='missing',pdf_path=NULL,pdf_note=?,updated_at=?,report_written=0 WHERE run_id=? AND candidate_hash=?",
               (args.reason.strip(), ledger.iso(ledger.now()), args.run, digest))
    path = workbook(db, args.output_root, row["started_at"], profile_path=run["profile_path"])
    db.execute("UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?", (args.run, digest))
    output(state="pdf_missing_recorded", report=str(path))


def lookup_name(args):
    if not getattr(args, "enable_local_excel_dedup", coworker_progress.LOCAL_EXCEL_DEDUP_ENABLED):
        output(state="disabled", enabled=False, duplicate=False,
               candidate_name=args.name.strip(), match_count=0, matches=[])
        return
    name = re.sub(r"\s+", " ", args.name.strip()).casefold()
    if not name:
        raise ValueError("candidate name required")
    if not args.db.is_file():
        raise ValueError("candidate library database not found")
    uri = args.db.resolve().as_uri() + "?mode=ro"
    with contextlib.closing(sqlite3.connect(uri, uri=True)) as db:
        db.row_factory = sqlite3.Row
        if not db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='assessments'").fetchone():
            raise ValueError("candidate library assessments table not found")
        rows = db.execute(
            "SELECT run_id,candidate_name,decision,send_status FROM assessments"
        ).fetchall()
    matches = [
        {"run_id": row["run_id"], "candidate_name": row["candidate_name"],
         "decision": row["decision"], "send_status": row["send_status"]}
        for row in rows
        if re.sub(r"\s+", " ", str(row["candidate_name"]).strip()).casefold() == name
    ]
    output(state="matches" if matches else "no_match", duplicate=bool(matches),
           candidate_name=args.name.strip(), match_count=len(matches), matches=matches)


def record_coworker_skip(db, args):
    observation = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
    result = coworker_progress.evaluate(observation, ledger.now())
    if not result["skip"]:
        raise ValueError("coworker_skip_requires_first_greeting_older_than_7_days")
    active = ledger.active_progress(db, args.run)
    digest = ledger.run_candidate_hash(db, args.run, args.candidate_id)
    if not active or active["candidate_hash"] != digest:
        raise ValueError("candidate_not_active")
    # The schema keeps an integer sentinel; both report sheets display this as unscored.
    record(db, args, payload={
        "candidate_id": args.candidate_id, "candidate_name": args.candidate_name,
        "started_at": active["started_at"], "overall_score": 0,
        "hard_requirements": False, "decision": "skip", "electronic_resume_enabled": False,
        "pros": "同事去重跳过，未作岗位评分",
        "cons": "首条同事招呼早于7天：" + str(result["first_greeting_at"]),
        "coworker_progress": observation,
    })


def rebuild(db, args):
    if args.day:
        day = date.fromisoformat(args.day)
        start = datetime.combine(day, time.min, CN_TZ).astimezone(timezone.utc)
        end = start + timedelta(days=1)
        rows = db.execute(
            "SELECT a.*,r.profile_path FROM assessments a JOIN runs r ON r.run_id=a.run_id "
            "WHERE a.started_at>=? AND a.started_at<? ORDER BY a.started_at",
            (start.isoformat(), end.isoformat()),
        ).fetchall()
    else:
        rows = db.execute(
            "SELECT a.*,r.profile_path FROM assessments a JOIN runs r ON r.run_id=a.run_id "
            "WHERE a.report_written=0 ORDER BY a.started_at"
        ).fetchall()
    if not rows:
        output(state="nothing_to_rebuild", day=args.day)
        return
    groups = {}
    for row in rows:
        key = (str(daily_dir(args.output_root, row["started_at"])), profile_identity(row["profile_path"]))
        groups.setdefault(key, []).append(row)
    paths = []
    for group in groups.values():
        row = group[0]
        path = workbook(db, args.output_root, row["started_at"], filename=args.output_file,
                        profile_path=row["profile_path"])
        paths.append(str(path))
        if not args.output_file:
            for item in group:
                db.execute("UPDATE assessments SET report_written=1 WHERE run_id=? AND candidate_hash=?",
                           (item["run_id"], item["candidate_hash"]))
    output(state="rebuilt", report=paths[0], reports=paths)


def message_summary(db,row):
    attempts=db.execute("SELECT message_index,status FROM message_attempts WHERE run_id=? AND candidate_hash=? ORDER BY message_index",(row["run_id"],row["candidate_hash"])).fetchall()
    return "；".join(f"第{a[0]}句:{a[1]}" for a in attempts) or row["send_status"]

def archive_card_files(root,started_at,name,digest,payload):
    from PIL import Image
    evidence=payload.get("card_evidence") or {}
    image=Path(evidence.get("image","")); uia=Path(evidence.get("uia_json",""))
    if not image.is_file() or not uia.is_file(): raise ValueError("card_evidence_required")
    with Image.open(image) as im:
        if min(im.size)<20: raise ValueError("card_screenshot_invalid")
        im.verify()
    observation=json.loads(uia.read_text(encoding="utf-8-sig"))
    if observation.get("name")!=name or observation.get("card")!=payload.get("card"):
        raise ValueError("card_evidence_identity_or_fields_mismatch")
    if not observation.get("observed_at") or not observation.get("bounds") or not observation.get("document_url","").startswith("https://www.zhipin.com/"):
        raise ValueError("card_uia_observation_incomplete")
    folder=daily_dir(root,started_at)/"名片图片"; folder.mkdir(parents=True,exist_ok=True)
    stem=re.sub(r'[<>:"/\\|?*\x00-\x1f]',"_",name)[:40]+"-"+digest[:12]
    paths=[]
    for source,suffix in ((image,".png"),(uia,".uia.json")):
        target=folder/(stem+suffix)
        if target.exists() and target.read_bytes()!=source.read_bytes(): raise ValueError("card_archive_collision")
        if source.resolve()!=target.resolve(): shutil.copyfile(source,target)
        paths.append(str(target))
    return paths

for _mutation_name in ("record","sync","archive","archive_resume","missing_resume","missing","finalize_timing","coworker_skip"):
    if _mutation_name in globals(): globals()[_mutation_name]=run_safety.fenced(globals()[_mutation_name])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=ledger.DEFAULT_DB)
    parser.add_argument("--output-root", type=Path, default=Path.home() / "Desktop")
    sub = parser.add_subparsers(dest="command", required=True)
    for name in ("record", "coworker-skip", "finalize-timing", "sync", "archive", "archive-resume", "missing-pdf", "missing-resume", "rebuild", "lookup-name"):
        command = sub.add_parser(name)
        if name == "lookup-name":
            command.add_argument("--name", required=True, help="Exact displayed candidate name to check in the local library")
            command.add_argument("--enable-local-excel-dedup", action="store_true",
                                 help="Explicitly enable optional local name deduplication (default: off)")
            continue
        if name != "rebuild":
            command.add_argument("--run", required=True)
        if name in ("record", "coworker-skip"):
            command.add_argument("--input", required=True, type=Path)
            if name == "coworker-skip":
                command.add_argument("--candidate-id", required=True)
                command.add_argument("--candidate-name", required=True)
        elif name == "finalize-timing":
            command.add_argument("--candidate-id", required=True)
            command.add_argument("--ended-at", required=True)
        elif name == "rebuild":
            command.add_argument("--day", help="Rebuild one local day (YYYY-MM-DD), including completed reports")
            command.add_argument("--output-file", help="Write a separate workbook filename in the daily folder")
        elif name != "rebuild":
            if name == "missing-resume":
                identity = command.add_mutually_exclusive_group(required=True)
                identity.add_argument("--candidate-id")
                identity.add_argument("--candidate-hash", help="Exact 64-character ledger candidate hash")
            else:
                command.add_argument("--candidate-id", required=True)
        if name == "archive":
            command.add_argument("--pdf", required=True)
        elif name == "archive-resume":
            command.add_argument("--image", required=True)
            command.add_argument("--ocr-text", required=True)
            command.add_argument("--ocr-json")
        elif name in ("missing-pdf", "missing-resume"):
            command.add_argument("--reason", required=True)
    args = parser.parse_args()
    try:
        if args.command == "lookup-name":
            lookup_name(args)
            return
        if args.command == "missing-resume" and args.candidate_hash and not re.fullmatch(r"[0-9a-fA-F]{64}", args.candidate_hash):
            raise ValueError("candidate hash must be exactly 64 hexadecimal characters")
        with ledger.connect(args.db) as db:
            table(db)
            {
                "record": record,
                "coworker-skip": record_coworker_skip,
                "finalize-timing": finalize_timing,
                "sync": sync,
                "archive": archive,
                "archive-resume": archive_resume,
                "missing-pdf": missing,
                "missing-resume": missing_resume,
                "rebuild": rebuild,
            }[args.command](db, args)
    except (OSError, ValueError, sqlite3.Error) as exc:
        output(state="error", reason=str(exc))
        sys.exit(2)


if __name__ == "__main__":
    main()
