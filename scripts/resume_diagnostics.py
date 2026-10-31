"""Keep resume failure reasons in the ledger and the visible run-status worksheet."""
from datetime import datetime, timedelta, timezone
import functools
import json
import sqlite3
import subprocess


def reason(error):
    text = str(error)
    if isinstance(error, subprocess.CalledProcessError):
        detail = error.stderr or error.stdout or ''
        if isinstance(detail, bytes):
            detail = detail.decode('utf-8-sig', errors='replace')
        text = f'子进程退出码 {error.returncode}；{detail.strip() or "未返回错误输出"}'
    elif isinstance(error, PermissionError):
        text = '没有写入权限，或文件被占用：' + text
    elif isinstance(error, FileNotFoundError):
        text = '所需程序或文件不存在：' + text
    return (type(error).__name__ + '：' + text)[-3000:]


def record(db, run_id, candidate_id, stage, error):
    if not isinstance(db, sqlite3.Connection):
        return  # Offline capture fixtures can supply a non-database stub.
    import ledger
    digest = ledger.run_candidate_hash(db, run_id, candidate_id)
    detail = json.dumps(dict(candidate_id=candidate_id, candidate_hash=digest, stage=stage,
                             reason=reason(error) if isinstance(error, BaseException) else str(error)),
                        ensure_ascii=False)
    if not db.execute("SELECT 1 FROM run_audit WHERE run_id=? AND event='resume_save_failed' AND detail=?",
                      (run_id, detail)).fetchone():
        db.execute('INSERT INTO run_audit(run_id,at,event,detail) VALUES (?,?,?,?)',
                   (run_id, ledger.iso(ledger.now()), 'resume_save_failed', detail))


def append_sheet(wb, db, profile, at):
    from report_paths import profile_identity
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet('运行状态')
    ws.append(['时间', '运行编号', '候选人', '处理阶段', '状态', '原因'])
    zone = timezone(timedelta(hours=8))
    day = datetime.fromisoformat(at).astimezone(zone).date()
    events = db.execute('''SELECT a.*,r.profile_path FROM run_audit a JOIN runs r ON r.run_id=a.run_id
        WHERE event='resume_save_failed' ORDER BY a.id''')
    for event in events:
        if profile_identity(event['profile_path']) != profile_identity(profile):
            continue
        data = json.loads(event['detail'])
        row = db.execute('SELECT candidate_name,started_at FROM assessments WHERE run_id=? AND candidate_hash=?',
                         (event['run_id'], data['candidate_hash'])).fetchone()
        progress = db.execute('SELECT started_at FROM candidate_progress WHERE run_id=? AND candidate_hash=?',
                              (event['run_id'], data['candidate_hash'])).fetchone()
        when = row['started_at'] if row else progress[0] if progress else event['at']
        if datetime.fromisoformat(when).astimezone(zone).date() != day:
            continue
        ws.append([datetime.fromisoformat(event['at']).astimezone(zone).strftime('%Y-%m-%d %H:%M:%S'),
                   event['run_id'], row['candidate_name'] if row else data['candidate_id'],
                   data['stage'], '保存失败（已记录）', data['reason']])
    # Include manual missing-file reasons and older records lacking structured events.
    import daily_metrics
    for row in daily_metrics.rows(db, profile, at):
        if row['evaluation_mode'] == 'card':
            continue
        payload=json.loads(row['evaluation_input'] or '{}')
        if not row['resume_image_path'] and row['resume_archive_status']!='missing':
            trigger=json.loads(row['eligibility_json'] or '{}').get('trigger')
            if trigger not in ('coworker_skip','coworker_unknown'):
                note=payload.get('resume_save_note')
                ws.append([datetime.fromisoformat(row['updated_at']).astimezone(zone).strftime('%Y-%m-%d %H:%M:%S'),row['run_id'],row['candidate_name'],'电子简历图片',
                           '未采集' if note else '待保存/未记录',
                           note or '本次没有已归档的简历图片，尚未记录具体保存结果；不代表已保存成功。'])
        for field, note, stage in (('resume_archive_status', 'resume_archive_note', '电子简历归档'),
                                   ('pdf_status', 'pdf_note', 'PDF归档')):
            if row[field] == 'missing' and row[note]:
                ws.append([datetime.fromisoformat(row['updated_at']).astimezone(zone).strftime('%Y-%m-%d %H:%M:%S'), row['run_id'], row['candidate_name'], stage,
                           '未保存', row[note]])
    if ws.max_row == 1:
        ws.append(['', '', '', '电子简历保存', '暂无失败记录', ''])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='26496A')
    for row in ws.iter_rows(min_row=2):
        import math
        lines=max(1,math.ceil(sum(2 if ord(c)>255 else 1 for c in str(row[-1].value or ''))/90))
        ws.row_dimensions[row[0].row].height=min(240,max(30,lines*16))
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value,str):cell.data_type='s'
    for col, width in zip('ABCDEF', (23, 22, 18, 24, 23, 95)):
        ws.column_dimensions[col].width = width


def archive_failure(stage, status_field):
    def decorate(fn):
        @functools.wraps(fn)
        def wrapped(db, args, *pos, **kwargs):
            try:
                return fn(db, args, *pos, **kwargs)
            except (OSError, ValueError) as exc:
                import ledger
                import candidate_log
                digest = ledger.run_candidate_hash(db, args.run, args.candidate_id)
                row = db.execute('SELECT * FROM assessments WHERE run_id=? AND candidate_hash=?',
                                 (args.run, digest)).fetchone()
                # A workbook-write failure after saving is not a resume-save failure.
                if not row or row[status_field] != 'saved':
                    record(db, args.run, args.candidate_id, stage, exc)
                    if row:
                        note = 'pdf_note' if status_field == 'pdf_status' else 'resume_archive_note'
                        db.execute(f"UPDATE assessments SET {status_field}='missing',{note}=?,report_written=0 WHERE run_id=? AND candidate_hash=?",
                                   (reason(exc), args.run, digest))
                    run = ledger.run_row(db, args.run)
                    try:
                        candidate_log.workbook(db, args.output_root, row['started_at'] if row else run['started_at'],
                                               profile_path=run['profile_path'])
                    except (OSError, ValueError):
                        pass  # The persisted reason will be included on the next export.
                raise
        return wrapped
    return decorate
