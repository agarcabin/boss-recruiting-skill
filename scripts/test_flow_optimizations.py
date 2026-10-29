"""Six workflow optimizations, all tested without real BOSS actions."""
import contextlib,io,json,os,tempfile,unittest
from pathlib import Path
from datetime import datetime,timezone,timedelta
from argparse import Namespace
from unittest.mock import Mock,patch
import candidate_log,candidate_checks,daily_metrics,fast_assessment,ledger,quick_start,recruiting_rules
import report_batch,run_targets,startup_plan,workflow_runner
import test_relaxed_controls as controls
call=controls.call


class FakeAdapter:
    def __init__(self,candidates,mode='resume',send_states=None,skip=False):
        self.candidates=list(candidates);self.events=[];self.send_states=list(send_states or ['sent']);self.skip=skip;self.read_count=0
    def select(self,processed):
        self.events.append('select')
        return self.candidates.pop(0) if self.candidates else None
    def open(self,candidate,mode):self.events.append('open')
    def colleague(self,candidate,mode):
        self.events.append('coworker')
        if self.skip:return dict(state='first_record',first_greeting_at=ledger.iso(ledger.now()-timedelta(days=9)),evidence='UIA同事首条实际时间')
        if mode=='card':return dict(state='absent',lookup_completed=True,evidence='实际卡片无同事图标')
        return dict(state='empty',evidence='UIA同事面板实际为空')
    def read(self,candidate,mode):
        self.events.append('read');self.read_count+=1
        if mode=='card':return dict(card=dict(advantage='组装',expectation='普工',already_contacted=False,coworker_contact=dict(state='absent',lookup_completed=True,evidence='实际卡片无同事图标')))
        return dict(work_experience=[dict(company='迈瑞医疗',duties='销售业务')],related_experience_years=2)
    def send(self,candidate,mode):
        self.events.append('send');return dict(state=self.send_states.pop(0) if self.send_states else 'sent')
    def verify(self,candidate,mode):
        self.events.append('verify');return dict(state='not-sent',may_retry=True)
    def favorite(self,candidate,mode):self.events.append('favorite');return dict(state='already_saved')
    def next(self,candidate,mode):self.events.append('next')


class FlowOptimizationTests(controls.RelaxedControlsTests):
    def policy(self,mode='resume'):
        rules=json.loads(recruiting_rules.CONFIG.read_text(encoding='utf-8'))
        rules.update(policy='meikang' if mode=='resume' else 'worker',mode=mode,greeting_mode='single')
        self.db.execute('UPDATE runs SET rule_snapshot=?,evaluation_mode=? WHERE run_id=?',(json.dumps(rules),mode,self.run))
        return rules

    def candidate(self,suffix='one'):
        return dict(candidate_id='candidate-'+suffix,candidate_name='测试人'+suffix,card_runtime='runtime-'+suffix)

    def test_record_retry_updates_existing_row_and_finished_visit(self):
        self.record();self.record(pros='更新后的实际依据')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidate_progress').fetchone()[0],1)
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM assessments').fetchone()[0],1)
        call(ledger.candidate_finish,self.db,run=self.run,candidate_id='candidate-one')
        self.record(pros='补写依据')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidate_progress').fetchone()[0],1)

    def test_pending_reservation_is_reused_and_unknown_retry_is_same_visit(self):
        self.record();first=self.reserve();second=self.reserve()
        self.assertEqual(first['token'],second['token'])
        self.resolve(first['token'],'unknown');third=self.reserve()
        self.assertNotEqual(first['token'],third['token'])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidate_progress').fetchone()[0],1)
        self.assertEqual(self.db.execute("SELECT status FROM message_attempt_history WHERE token=?",(first['token'],)).fetchone()[0],'unknown')

    def test_note_retry_and_check_supplement_do_not_create_visits(self):
        self.record(decision='review')
        call(ledger.note,self.db,run=self.run,candidate_id='candidate-one',result='review')
        call(ledger.candidate_finish,self.db,run=self.run,candidate_id='candidate-one')
        call(ledger.note,self.db,run=self.run,candidate_id='candidate-one',result='review')
        candidate_checks.record(self.db,self.run,'candidate-one',dict(state='empty',evidence='当前实际同事面板'))
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidate_progress').fetchone()[0],1)

    def test_true_reopening_is_still_allowed_and_creates_new_visit(self):
        self.record();self.resolve(self.reserve()['token'])
        call(ledger.candidate_begin,self.db,run=self.run,candidate_id='candidate-one',new_visit=True)
        self.record();self.resolve(self.reserve()['token'])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidate_progress').fetchone()[0],2)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM candidates WHERE status='sent'").fetchone()[0],2)

    def test_late_unknown_attempt_confirmation_is_retained(self):
        self.record();first=self.reserve();self.resolve(first['token'],'unknown')
        second=self.reserve();self.resolve(first['token'],'sent');self.resolve(second['token'],'unknown')
        self.assertEqual(self.db.execute('SELECT status FROM candidates').fetchone()[0],'sent')
        digest=ledger.run_candidate_hash(self.db,self.run,'candidate-one')
        self.assertEqual(len(candidate_log.attempt_rows(self.db,self.run,digest)),2)

    def test_batch_exports_once_and_marks_written_only_after_export(self):
        original=candidate_log.workbook;exports=[]
        def counted(*a,**kw):
            if not report_batch.active(a[0]):exports.append(True)
            return original(*a,**kw)
        with patch.object(candidate_log,'workbook',side_effect=counted):
            with report_batch.batch(self.db):
                self.record();self.resolve(self.reserve()['token'])
                call(candidate_log.sync,self.db,run=self.run,candidate_id='candidate-one',output_root=self.root)
                call(candidate_log.finalize_timing,self.db,run=self.run,candidate_id='candidate-one',ended_at=ledger.iso(ledger.now()),output_root=self.root)
                self.assertEqual(self.db.execute('SELECT report_written FROM assessments').fetchone()[0],0)
        self.assertEqual(len(exports),1)
        self.assertEqual(self.db.execute('SELECT report_written FROM assessments').fetchone()[0],1)

    def test_batch_export_failure_is_nonblocking(self):
        original=candidate_log.workbook
        def blocked(*a,**kw):
            if not report_batch.active(a[0]):raise PermissionError('locked')
            return original(*a,**kw)
        with patch.object(candidate_log,'workbook',side_effect=blocked):
            with report_batch.batch(self.db) as batch:self.record()
        self.assertTrue(batch['warnings'])
        self.assertEqual(self.db.execute('SELECT report_written FROM assessments').fetchone()[0],0)
        self.assertEqual(self.reserve()['state'],'ready')

    def test_cached_workbook_reads_only_when_changed(self):
        result=self.record();path=Path(result['report']);daily_metrics._workbook_cache.clear()
        from openpyxl import load_workbook
        with patch('openpyxl.load_workbook',wraps=load_workbook) as reader:
            daily_metrics.workbook_counts(path);daily_metrics.workbook_counts(path)
            self.assertEqual(reader.call_count,1)
            stat=path.stat();os.utime(path,ns=(stat.st_atime_ns,stat.st_mtime_ns+1000000))
            daily_metrics.workbook_counts(path)
            self.assertEqual(reader.call_count,2)

    def test_direct_match_needs_no_full_scorecard(self):
        self.policy()
        payload=dict(candidate_id='candidate-one',candidate_name='测试人',coworker_progress=dict(state='empty',evidence='当前同事无记录'),work_experience=[dict(company='迈瑞医疗',duties='销售')],related_experience_years=2)
        with patch.object(recruiting_rules,'evaluate',side_effect=AssertionError('full evaluation must be skipped')):
            result=call(lambda db,args:candidate_log.record(db,args,payload),self.db,run=self.run,output_root=self.root)
        row=self.db.execute('SELECT * FROM assessments').fetchone()
        self.assertEqual(result['decision'],'greet');self.assertEqual(row['overall_score'],75)
        self.assertEqual(json.loads(row['eligibility_json'])['method'],'direct_rule')

    def test_direct_match_still_checks_colleague_first(self):
        self.policy()
        with patch.object(fast_assessment,'evaluate',side_effect=AssertionError('must not match before coworker')):
            with self.assertRaisesRegex(ValueError,'coworker_check_before_analysis'):
                self.record(coworker_progress=dict(state='unavailable',evidence='面板不可读'))

    def test_short_direct_match_skips_long_ocr(self):
        self.policy()
        from workflow_windows import WindowsAdapter
        adapter=WindowsAdapter(self.db,self.run,self.root,hwnd=1)
        adapter.query=Mock(return_value=dict(keyword_text='经历概览\n迈瑞医疗\n销售经理',source='actual_uia_overview'))
        with patch('workflow_windows.capture_resume.capture',side_effect=AssertionError('must not OCR a direct match')):
            facts=adapter.read(self.candidate(),'resume')
        self.assertEqual(facts['read_source'],'actual_uia_overview')

    def test_target_defaults_are_linked_but_explicit_view_goal_is_honored(self):
        result=self.start(target=500)
        self.assertEqual((result['target'],result['max_views']),(500,500))
        self.assertEqual(self.start(target=500,max_views=100)['max_views'],100)

    def test_runner_processes_two_candidates_and_stops_without_opening_third(self):
        self.run=self.start(target=2,max_views=10)['run_id'];self.policy()
        adapter=FakeAdapter([self.candidate('one'),self.candidate('two'),self.candidate('three')])
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        first=runner.advance();second=runner.advance()
        self.assertEqual(first['state'],'candidate_complete');self.assertEqual(second['state'],'complete')
        self.assertEqual(adapter.events.count('open'),2);self.assertEqual(adapter.events.count('send'),2)
        self.assertEqual(adapter.events.count('next'),1)
        self.assertEqual(len(second['exports']),1)

    def test_runner_exports_before_switching_to_next_candidate(self):
        self.policy();adapter=FakeAdapter([self.candidate()])
        def next_candidate(candidate,mode):
            from report_paths import report_path
            path=report_path(self.root,ledger.iso(ledger.now()),self.profile)
            self.assertTrue(path.is_file())
            self.assertEqual(daily_metrics.workbook_counts(path)['sent'],1)
        adapter.next=next_candidate
        self.assertEqual(workflow_runner.Runner(self.db,self.run,adapter,self.root).advance()['state'],'candidate_complete')

    def test_runner_worker_mode_uses_same_pipeline(self):
        self.policy('card');adapter=FakeAdapter([self.candidate()],mode='card')
        result=workflow_runner.Runner(self.db,self.run,adapter,self.root).advance()
        self.assertEqual(result['state'],'candidate_complete')
        self.assertEqual(adapter.events.count('send'),1)
        self.assertLess(adapter.events.index('coworker'),adapter.events.index('read'))

    def test_runner_waiting_status_does_not_rewrite_checkpoint(self):
        self.policy();adapter=FakeAdapter([self.candidate()]);adapter.read=Mock(return_value=dict(keyword_text='实际普通销售经历'))
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root);runner.advance()
        self.db.execute('CREATE TABLE writes(value INTEGER)')
        self.db.execute("CREATE TRIGGER count_workflow_updates AFTER UPDATE ON workflow_state BEGIN INSERT INTO writes VALUES(1); END")
        runner.advance();runner.advance()
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM writes').fetchone()[0],0)

    def test_native_adapter_keeps_uia_working_when_card_is_occluded(self):
        from workflow_windows import WindowsAdapter
        adapter=WindowsAdapter(self.db,self.run,self.root,hwnd=1)
        observation=dict(card=dict(coworker_contact=dict(state='absent',lookup_completed=True,evidence='卡片UIA')))
        adapter.ps=Mock(side_effect=[RuntimeError('boss_window_not_foreground'),observation])
        result=adapter.colleague(self.candidate(),'card')
        self.assertEqual(result['state'],'absent')
        self.assertIn('-NoImage',adapter.ps.call_args.args)

    def test_runner_skips_colleague_without_reading_or_sending(self):
        self.policy();adapter=FakeAdapter([self.candidate()],skip=True)
        result=workflow_runner.Runner(self.db,self.run,adapter,self.root).advance()
        self.assertEqual(result['state'],'candidate_complete')
        self.assertNotIn('read',adapter.events);self.assertNotIn('send',adapter.events);self.assertNotIn('favorite',adapter.events)

    def test_runner_retries_unknown_once_without_extra_view(self):
        self.policy();adapter=FakeAdapter([self.candidate()],send_states=['unknown','sent'])
        result=workflow_runner.Runner(self.db,self.run,adapter,self.root).advance()
        self.assertEqual(result['state'],'candidate_complete')
        self.assertEqual(adapter.events.count('send'),2)
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidate_progress').fetchone()[0],1)
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM assessments').fetchone()[0],1)

    def test_runner_pause_and_resume_from_same_phase(self):
        self.policy();adapter=FakeAdapter([self.candidate()])
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        call(ledger.control,self.db,run=self.run,action='pause')
        self.assertEqual(runner.advance()['state'],'paused');self.assertEqual(adapter.events,[])
        call(ledger.control,self.db,run=self.run,action='resume')
        self.assertEqual(runner.advance()['state'],'candidate_complete')

    def test_runner_handoff_does_not_repeat_reads_or_full_evaluation(self):
        self.policy();adapter=FakeAdapter([self.candidate()])
        adapter.read=Mock(return_value=dict(keyword_text='普通行业海外销售经历'))
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        first=runner.advance();self.assertEqual(first['state'],'needs_assessment')
        with patch.object(fast_assessment,'evaluate',side_effect=AssertionError('waiting should not re-evaluate')):
            self.assertEqual(runner.advance()['state'],'needs_assessment')
        adapter.read.assert_called_once()
        Path(first['response_file']).write_text(json.dumps(dict(overall_score=80,hard_requirements=True,decision='greet',electronic_resume_enabled=True,electronic=dict(score=80,evidence='实际已读海外销售'),pros='已读经历',cons='未知项')),encoding='utf-8')
        self.assertEqual(runner.advance()['state'],'candidate_complete')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidate_progress').fetchone()[0],1)

    def test_status_reader_reports_waiting_assessment_after_reload(self):
        self.policy();adapter=FakeAdapter([self.candidate()]);adapter.read=Mock(return_value=dict(keyword_text='实际其他销售经历'))
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root);runner.advance()
        reader=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        self.assertEqual(reader.status()['state'],'needs_assessment')
        self.assertTrue(reader.status()['request_file'])

    def test_status_reader_reports_page_unavailable_after_worker_exit(self):
        self.policy();adapter=FakeAdapter([])
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        self.assertEqual(runner.advance()['state'],'needs_page')
        reader=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        self.assertEqual(reader.status()['state'],'needs_page')

    def test_runner_restart_after_dispatch_verifies_instead_of_resending(self):
        self.policy();adapter=FakeAdapter([self.candidate()])
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        while runner.phase!='send':runner.step();runner.save()
        reservation=self.reserve();runner.data['token']=reservation['token'];runner.data['attempts']=1
        runner.phase='verify';runner.save()
        adapter.verify=Mock(return_value=dict(state='sent'))
        resumed=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        self.assertEqual(resumed.advance()['state'],'candidate_complete')
        self.assertNotIn('send',adapter.events)


class StartupTests(unittest.TestCase):
    def test_target_normalization(self):
        self.assertEqual(run_targets.resolve(),(150,300));self.assertEqual(run_targets.resolve(500),(500,500))
        self.assertEqual(run_targets.resolve(500,100),(500,100))

    def test_saved_plan_resumes_without_reparsing_or_duplicate_start(self):
        from test_ledger import PROFILE
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);profile=root/'profile.md';profile.write_text(PROFILE,encoding='utf-8')
            projects=[dict(profile=str(profile),company='美康',position='国际销售经理',approval='approved')]
            with patch.object(quick_start.project_registry,'list_projects',return_value=projects),patch.object(ledger,'PROFILE_ROOT',root):
                def main(args):
                    import sys
                    out=io.StringIO()
                    with patch.object(sys,'argv',['quick_start.py']+args),contextlib.redirect_stdout(out):quick_start.main()
                    return json.loads(out.getvalue())
                plan=main(['--prompt','帮我打招呼，打美康500人','--state-dir',str(root/'plans')])
                with patch.object(quick_start,'parse',side_effect=AssertionError('must not parse again')):
                    resumed=main(['--plan',plan['plan_file'],'--db',str(root/'ledger.db'),'--confirm-filters'])
                    again=main(['--plan',plan['plan_file'],'--db',str(root/'ledger.db'),'--confirm-filters'])
                self.assertEqual(resumed['run']['run_id'],again['run']['run_id'])
                self.assertEqual(resumed['max_views'],500)

    def test_startup_actions_are_acknowledged_once(self):
        with tempfile.TemporaryDirectory() as directory:
            plan=startup_plan.save(dict(state='await_filters',next_actions=['sidebar']), 'original',directory)
            startup_plan.complete(plan,'sidebar');startup_plan.complete(plan,'sidebar')
            loaded=startup_plan.load(plan['plan_file'])
            self.assertEqual(loaded['startup_completed'],['sidebar'])


if __name__=='__main__':unittest.main()
