"""Offline acceptance tests for shared pacing, resume diagnostics and update routes."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import timedelta
import json
from pathlib import Path
import subprocess
import unittest
from unittest.mock import Mock, patch
from urllib.error import URLError

from openpyxl import load_workbook
import candidate_log
import capture_resume
import greeting_rate
import ledger
import progress_window
import run_safety
import resume_diagnostics
import skill_update
import update_transport
import workflow_runner
import test_flow_optimizations as flow_tests
import test_relaxed_controls as control_tests
import test_skill_update as updater_tests
call=control_tests.call


class PacingAndDiagnosticsTests(unittest.TestCase):
    setUp = control_tests.RelaxedControlsTests.setUp
    start = control_tests.RelaxedControlsTests.start
    payload = control_tests.RelaxedControlsTests.payload
    record = control_tests.RelaxedControlsTests.record
    reserve = control_tests.RelaxedControlsTests.reserve
    resolve = control_tests.RelaxedControlsTests.resolve
    policy = flow_tests.FlowOptimizationTests.policy

    def test_fast_machine_waits_only_for_remainder_and_resumes_at_boundary(self):
        begin = ledger.now()
        for index in range(10):
            self.record('candidate-'+str(index))
            with patch.object(ledger, 'now', return_value=begin+timedelta(seconds=index*5)):
                self.assertEqual(self.reserve('candidate-'+str(index))['state'], 'ready')
        self.record('eleventh')
        with patch.object(ledger, 'now', return_value=begin+timedelta(seconds=60)):
            result = self.reserve('eleventh')
            self.assertEqual((result['state'],result['retry_after_seconds']), ('rate_limited',240))
            self.assertEqual(self.db.execute('SELECT COUNT(*) FROM candidates').fetchone()[0],10)
        with patch.object(ledger, 'now', return_value=begin+timedelta(seconds=299)):
            self.assertEqual(self.reserve('eleventh')['retry_after_seconds'],1)
        with patch.object(ledger, 'now', return_value=begin+timedelta(seconds=300)):
            self.assertEqual(self.reserve('eleventh')['state'],'ready')

    def test_slow_machine_has_no_extra_wait_or_accumulated_burst(self):
        begin = ledger.now()
        for index in range(12):
            self.record('candidate-'+str(index))
            with patch.object(ledger, 'now', return_value=begin+timedelta(seconds=index*35)):
                self.assertEqual(self.reserve('candidate-'+str(index))['state'],'ready')
        with patch.object(ledger, 'now', return_value=begin+timedelta(hours=2)):
            for index in range(10):
                self.assertEqual(greeting_rate.claim(self.db,'new-'+str(index))['state'],'ready')
            self.assertEqual(greeting_rate.claim(self.db,'new-11')['state'],'rate_limited')

    def test_pending_unknown_retry_and_followup_do_not_add_people(self):
        self.record()
        first = self.reserve()
        self.assertEqual(self.reserve()['token'],first['token'])
        self.resolve(first['token'],'unknown')
        self.assertEqual(self.reserve()['state'],'ready')
        self.assertEqual(greeting_rate.status(self.db)['used'],1)

    def test_switching_job_run_and_reopening_database_preserves_limit(self):
        for index in range(10):
            greeting_rate.claim(self.db,str(index))
        self.run = self.start()['run_id']
        self.policy('card')
        self.record('new-job',electronic_resume_enabled=False,electronic=None,
                    card=dict(advantage='组装',expectation='普工',already_contacted=False,
                              coworker_contact=dict(state='absent',lookup_completed=True,evidence='当前卡片没有同事图标')))
        self.assertEqual(self.reserve('new-job')['state'],'rate_limited')
        with closing(ledger.connect(self.root/'ledger.db')) as other:
            self.assertEqual(greeting_rate.claim(other,'other-process')['state'],'rate_limited')

    def test_concurrent_process_connections_can_only_claim_ten_slots(self):
        path = self.root/'ledger.db'
        # Independent SQLite connections exercise the cross-process transaction boundary.
        def claim(index):
            with closing(ledger.connect(path)) as db:
                return greeting_rate.claim(db,'concurrent-'+str(index))['state']
        with ThreadPoolExecutor(max_workers=4) as pool:
            states=list(pool.map(claim,range(20)))
        self.assertEqual(states.count('ready'),10)
        self.assertEqual(states.count('rate_limited'),10)

    def test_scheduler_waits_then_sends_without_rereading_and_honors_pause(self):
        self.policy('card')
        for index in range(10):
            greeting_rate.claim(self.db,'fill-'+str(index))
        begin = ledger.now()
        adapter = flow_tests.FakeAdapter([dict(candidate_id='candidate-one',candidate_name='测试人',card_runtime='card-one')],mode='card')
        runner = workflow_runner.Runner(self.db,self.run,adapter,self.root)
        self.assertEqual(runner.advance()['state'],'rate_limited')
        self.assertNotIn('send',adapter.events)
        self.db.execute("UPDATE run_control SET state='paused' WHERE run_id=?",(self.run,))
        self.assertEqual(runner.advance()['state'],'paused')
        self.db.execute("UPDATE run_control SET state='running' WHERE run_id=?",(self.run,))
        # A new scheduler instance resumes the exact saved candidate/phase.
        runner = workflow_runner.Runner(self.db,self.run,adapter,self.root)
        with patch.object(ledger,'now',return_value=begin+timedelta(seconds=301)):
            runner.advance()
        self.assertEqual(adapter.events.count('read'),1)
        self.assertEqual(adapter.events.count('send'),1)

    def report_rows(self,path):
        wb = load_workbook(path)
        try:
            self.assertIn('运行状态',wb.sheetnames)
            self.assertEqual(wb['运行状态'].sheet_state,'visible')
            return list(wb['运行状态'].values)
        finally:
            wb.close()

    def test_countdown_is_visible_and_stopping_during_wait_prevents_send(self):
        self.policy('card')
        for index in range(10):greeting_rate.claim(self.db,'fill-'+str(index))
        adapter=flow_tests.FakeAdapter([dict(candidate_id='candidate-one',candidate_name='测试人',card_runtime='card-one')],mode='card')
        runner=workflow_runner.Runner(self.db,self.run,adapter,self.root)
        self.assertEqual(runner.advance()['state'],'rate_limited')
        self.assertIn('速度控制',progress_window.snapshot(self.db,self.run,self.root)['current_step'])
        run_safety.request_stop(self.db,self.run,'manual_stop')
        self.assertEqual(runner.advance()['state'],'stopped')
        self.assertNotIn('send',adapter.events)

    def test_intentionally_skipped_capture_has_explanation_without_false_failure(self):
        result=self.record(resume_save_note='关键词直接通道，未执行可选长图采集')
        rows=self.report_rows(result['report'])
        self.assertTrue(any('未采集' in str(row) and '关键词直接通道' in str(row) for row in rows))
        self.assertFalse(any('保存失败' in str(row) for row in rows))

    def test_read_warning_survives_sync_and_is_visible_with_candidate(self):
        result=self.record(read_warning='OCR模型缺失：rapidocr model not found')
        self.resolve(self.reserve()['token'])
        call(candidate_log.sync,self.db,run=self.run,candidate_id='candidate-one',output_root=self.root)
        rows=self.report_rows(result['report'])
        self.assertTrue(any('测试人' in str(row) and 'OCR模型缺失' in str(row) for row in rows))
        self.assertEqual(self.db.execute('SELECT resume_archive_status FROM assessments').fetchone()[0],'missing')

    def test_archive_permission_error_is_written_without_blocking_record(self):
        with patch.object(candidate_log,'archive_resume_files',side_effect=PermissionError(13,'Access denied','D:/简历')):
            result=self.record(resume_evidence=dict(image='source.png',ocr_text='source.txt'))
        self.assertEqual(result['state'],'recorded')
        rows=self.report_rows(result['report'])
        self.assertTrue(any('Access denied' in str(row) and '电子简历归档' in str(row) for row in rows))

    def test_subprocess_stderr_is_kept_and_failure_before_assessment_still_exports(self):
        ledger.ensure_candidate(self.db,self.run,'before-assessment')
        error=subprocess.CalledProcessError(1,['local-capture'],stderr='c_resume_document_missing: 无法定位简历区域')
        with patch.object(capture_resume,'check_candidate'), patch.object(capture_resume.powershell_host,'command',return_value=['capture']), \
                patch.object(capture_resume.subprocess,'run',side_effect=error):
            with self.assertRaises(subprocess.CalledProcessError):
                capture_resume.capture(self.db,self.run,'before-assessment',123,self.root/'capture')
        run=ledger.run_row(self.db,self.run)
        path=candidate_log.workbook(self.db,self.root,run['started_at'],profile_path=self.profile)
        rows=self.report_rows(path)
        self.assertTrue(any('before-assessment' in str(row) and 'c_resume_document_missing' in str(row) for row in rows))
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM assessments').fetchone()[0],0)

    def test_explicit_archive_failure_is_recorded_and_reported(self):
        self.record()
        with self.assertRaisesRegex(ValueError,'resume_image_missing'):
            call(candidate_log.archive_resume,self.db,run=self.run,candidate_id='candidate-one',output_root=self.root,
                 image=str(self.root/'missing.png'),ocr_text=str(self.root/'missing.txt'),ocr_json=None)
        run=ledger.run_row(self.db,self.run)
        path=candidate_log.workbook(self.db,self.root,run['started_at'],profile_path=self.profile)
        self.assertTrue(any('resume_image_missing' in str(row) for row in self.report_rows(path)))

    def test_manual_missing_reason_and_profile_isolation(self):
        self.record()
        call(candidate_log.missing_resume,self.db,run=self.run,candidate_id='candidate-one',candidate_hash=None,
             output_root=self.root,reason='截图被其他窗口遮挡，无法保存')
        run=ledger.run_row(self.db,self.run)
        path=candidate_log.workbook(self.db,self.root,run['started_at'],profile_path=self.profile)
        self.assertTrue(any('截图被其他窗口遮挡' in str(row) for row in self.report_rows(path)))


class UpdateRouteTests(unittest.TestCase):
    setUp=updater_tests.SkillUpdateTests.setUp
    fetch=updater_tests.SkillUpdateTests.fetch

    def test_metadata_falls_back_and_next_check_starts_with_successful_route(self):
        seen=[]
        good=update_transport.MIRRORS[1]
        def fetch(url,limit=skill_update.LIMIT):
            seen.append(url)
            if url.startswith(good):return json.dumps(self.release).encode()
            raise URLError('connection timed out')
        with patch.object(skill_update,'fetch',side_effect=fetch):
            first=skill_update.check(self.local)
            self.assertEqual(first['update_route'],good)
            self.assertEqual(len(first['route_failures']),2)
            seen.clear()
            self.assertEqual(skill_update.check(self.local)['state'],'update_available')
            self.assertEqual(len(seen),1)
            self.assertTrue(seen[0].startswith(good))

    def test_failed_preferred_route_falls_back_and_replaces_preference(self):
        update_transport.remember(self.local,'metadata',update_transport.MIRRORS[0])
        def fetch(url,limit=skill_update.LIMIT):
            if url.startswith(update_transport.MIRRORS[0]):raise URLError('offline')
            return json.dumps(self.release).encode()
        with patch.object(skill_update,'fetch',side_effect=fetch):
            self.assertEqual(skill_update.check(self.local)['update_route'],'direct')
        self.assertEqual(update_transport.memory(self.local)['metadata'],'direct')

    def test_download_only_mirror_can_find_latest_version_from_canonical_page(self):
        good=update_transport.MIRRORS[0]
        page=f'<html><meta property="og:url" content="https://github.com/{self.repository}/releases/tag/v1.01"></html>'.encode()
        def fetch(url,limit=skill_update.LIMIT):
            if url.startswith(good) and url.startswith(good+'https://github.com/'):
                return page
            raise URLError('GitHub API is blocked')
        with patch.object(skill_update,'fetch',side_effect=fetch):
            result=skill_update.check(self.local)
        self.assertEqual(result['latest_version'],'1.01')
        self.assertEqual(result['archive_url'],self.archive_url)

    def test_download_rejects_bad_checksum_then_remembers_validated_route(self):
        release=skill_update.release_result(skill_update.source(self.local),'1.0.0',self.release)
        good=update_transport.MIRRORS[1]
        def fetch(url,limit=skill_update.LIMIT):
            if url.startswith('https://github.com/'):raise URLError('direct download timeout')
            if url.endswith('.sha256') and url.startswith(update_transport.MIRRORS[0]):return b'0'*64
            return self.fetch(url,limit)
        with patch.object(skill_update,'fetch',side_effect=fetch):
            data,_=skill_update.download(self.local,release)
        self.assertEqual(data,self.archive)
        self.assertEqual(release['download_route'],good)
        self.assertEqual(update_transport.memory(self.local)['download'],good)
        self.assertEqual(len(release['download_route_failures']),2)

    def test_html_download_and_matching_but_invalid_zip_try_another_route(self):
        release=skill_update.release_result(skill_update.source(self.local),'1.0.0',self.release)
        bad=b'<html>service unavailable</html>'
        def fetch(url,limit=skill_update.LIMIT):
            if url.startswith('https://github.com/'):
                return (skill_update.digest(bad)+'  file.zip').encode() if url.endswith('.sha256') else bad
            return self.fetch(url,limit)
        with patch.object(skill_update,'fetch',side_effect=fetch):
            data,_=skill_update.download(self.local,release)
        self.assertEqual(data,self.archive)
        self.assertNotEqual(release['download_route'],'direct')

    def test_all_routes_fail_with_each_route_explained_and_installation_untouched(self):
        before=(self.local/'SKILL.md').read_bytes()
        with patch.object(skill_update,'fetch',side_effect=URLError('DNS resolution failed')):
            with self.assertRaisesRegex(ValueError,'all_update_routes_failed') as error:
                skill_update.update(self.local,self.base/'backups',self.db)
        for route in ('direct',*update_transport.MIRRORS):self.assertIn(route,str(error.exception))
        self.assertEqual(before,(self.local/'SKILL.md').read_bytes())

    def test_invalid_response_or_other_repository_page_does_not_poison_memory(self):
        with patch.object(skill_update,'fetch',return_value=b'<html>Login required</html>'):
            with self.assertRaisesRegex(ValueError,'all_update_routes_failed'):
                skill_update.check(self.local)
        self.assertEqual(update_transport.memory(self.local),{})
        with self.assertRaises(ValueError):
            update_transport.release_from_page(b'<link rel="canonical" href="https://github.com/other/repo/releases/tag/v1.01">',self.repository)

    def test_route_storage_corruption_is_recoverable_and_preferences_are_separate(self):
        file=update_transport.memory_file(self.local)
        file.parent.mkdir(parents=True)
        file.write_text('broken json',encoding='utf-8')
        update_transport.remember(self.local,'metadata',update_transport.MIRRORS[0])
        update_transport.remember(self.local,'download',update_transport.MIRRORS[2])
        self.assertEqual(update_transport.routes(self.local,'metadata')[0],update_transport.MIRRORS[0])
        self.assertEqual(update_transport.routes(self.local,'download')[0],update_transport.MIRRORS[2])


if __name__=='__main__':unittest.main()
