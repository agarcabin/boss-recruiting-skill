"""Create a local, non-sending readiness report for the recruiting skill."""

from __future__ import annotations

import argparse
import importlib.metadata
import json
import os
import platform
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path


try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass


def check(name, action):
    try:
        detail = action()
        return {"name": name, "status": "通过", "detail": str(detail)}
    except Exception as exc:
        return {"name": name, "status": "失败", "detail": str(exc)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, default=Path.home() / "Desktop")
    parser.add_argument("--browser-verified", action="store_true")
    parser.add_argument("--browser-evidence", default="")
    args = parser.parse_args()
    if args.browser_verified and len(args.browser_evidence.strip()) < 15:
        parser.error("--browser-verified requires specific --browser-evidence")
    skill = Path(__file__).resolve().parents[1]
    mouse = skill / "scripts" / "move-click.ps1"
    uia_query = skill / "scripts" / "uia-query.ps1"
    resume_capture = skill / "scripts" / "resume_capture.ps1"
    idle_mouse = skill / "scripts" / "idle_mouse.py"
    ledger = skill / "scripts" / "ledger.py"
    folder = args.output_root / "自动投递记录" / "标准化测试"
    checks = []
    checks.append(check("Python 与 Excel 库", lambda: __import__("openpyxl").__version__))
    checks.append(check("PDF 解析库", lambda: __import__("pypdf").__version__))
    checks.append(check("RapidOCR", lambda: importlib.metadata.version("rapidocr")))
    checks.append(check("ONNX Runtime", lambda: importlib.metadata.version("onnxruntime")))
    checks.append(check("OpenCV", lambda: __import__("cv2").__version__))

    def portable_runtime_test():
        python = skill / "runtime" / "win-x64" / "python" / "python.exe"
        models = skill / "runtime" / "win-x64" / "python" / "Lib" / "site-packages" / "rapidocr" / "models"
        required = [
            models / "PP-OCRv6_det_small.onnx",
            models / "PP-OCRv6_rec_small.onnx",
            models / "ch_ppocr_mobile_v2.0_cls_mobile.onnx",
        ]
        if not python.is_file():
            raise RuntimeError("便携 Python 缺失")
        missing = [path.name for path in required if not path.is_file()]
        if missing:
            raise RuntimeError(f"RapidOCR 模型缺失: {', '.join(missing)}")
        return f"{python}；当前测试解释器={sys.executable}"

    checks.append(check("便携运行时", portable_runtime_test))

    def window_test():
        import tkinter as tk
        root = tk.Tk()
        try:
            root.withdraw()
            root.update()
            return f"Tk {root.tk.call('info', 'patchlevel')} 可创建窗口"
        finally:
            root.destroy()

    checks.append(check("进度窗口", window_test))

    def project_package_test():
        import project_registry
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "profiles" / "source"
            source.mkdir(parents=True)
            profile = source / "profile.md"
            profile.write_text("""---
schema: boss-greeting-profile/v1
company: Test Company
position: Sales
approval: approved
approved_by: owner
approved_at: 2026-10-20
---
## 页面核验
Test.
## 企业与岗位事实
Test.
## 自动筛选条件
Test.
## 评分规则
Test.
## 加分项
Test.
## 待人工确认
None.
## 固定招呼话术
Approved test greeting.
## 冲突与待确认
无
## 来源
Test.
""", encoding="utf-8")
            project_registry.register_project(profile, "standard-test")
            archive = root / "standard-test.boss-project.zip"
            project_registry.export_project(profile, archive)
            imported = project_registry.import_project(archive, root / "imported")
            assert imported["project_id"] == "standard-test"
            assert project_registry.list_projects(root / "imported")[0]["company"] == "Test Company"
        return "项目登记、SHA-256 分享包和导入通过"

    checks.append(check("项目包导入导出", project_package_test))

    def desktop_test():
        folder.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=folder) as tmp:
            path = Path(tmp) / "test.xlsx"
            from openpyxl import Workbook, load_workbook
            wb = Workbook()
            wb.active["A1"] = "通过"
            wb.save(path)
            reopened = load_workbook(path, read_only=True)
            try:
                assert reopened.active["A1"].value == "通过"
            finally:
                reopened.close()
        return str(folder)

    checks.append(check("桌面目录与 XLSX 读写", desktop_test))

    def ledger_test():
        import importlib.util
        spec = importlib.util.spec_from_file_location("recruiting_ledger_test", ledger)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as tmp:
            db = module.connect(Path(tmp) / "ledger.sqlite3")
            assert db.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 0
            db.close()
        return "临时 SQLite 创建、读取成功"

    checks.append(check("招聘账本", ledger_test))

    def workflow_test():
        result = subprocess.run([sys.executable, "-m", "unittest", "discover", "-s", str(skill / "scripts"), "-p", "test_*.py"],
                                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
        if result.returncode:
            raise RuntimeError((result.stderr or result.stdout or "unittest failed")[-800:])
        return "候选人记录（含判断时长与绿黄行配色）、限额、恢复和简历/PDF 归档测试通过"

    checks.append(check("招聘流程测试", workflow_test))
    checks.append(check("PowerShell 7", lambda: shutil.which("pwsh") or (_ for _ in ()).throw(RuntimeError("pwsh 未安装或未在 PATH"))))

    def mouse_test():
        pwsh = shutil.which("pwsh")
        if not pwsh:
            raise RuntimeError("pwsh 不可用")
        if not mouse.is_file():
            raise RuntimeError("鼠标脚本不存在")
        source = mouse.read_text(encoding="utf-8")
        if "ShowWindow($WindowHandle, 9)" in source or "IsZoomed($WindowHandle)" not in source:
            raise RuntimeError("鼠标脚本可能还原最大化窗口")
        if "ForbiddenNames" not in source or "ForbiddenClassNames" not in source:
            raise RuntimeError("鼠标脚本缺少危险控件拦截")
        if "PostClickMoveMs = 1000" not in source or "Invoke-PostClickMove" not in source:
            raise RuntimeError("鼠标脚本缺少点击后 1 秒随机移动")
        if "clickStartedAt" not in source or "clickEndedAt" not in source:
            raise RuntimeError("鼠标脚本缺少点击时间边界输出")
        result = subprocess.run([pwsh, "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(mouse),
                                 "-X", "100", "-Y", "100", "-ContentBounds", "0,0,200,200", "-DryRun"],
                                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
        if result.returncode:
            raise RuntimeError((result.stderr or result.stdout or "mouse dry run failed")[-500:])
        data = json.loads(result.stdout)
        if not data.get("dryRun") or data.get("clicked"):
            raise RuntimeError("dry run 结果异常")
        if "clickStartedAt" not in data or "clickEndedAt" not in data or "postClickMove" not in data:
            raise RuntimeError("dry run 缺少点击时间边界或点击后移动字段")
        return "轨迹生成成功，未点击"

    checks.append(check("鼠标轨迹 dry run", mouse_test))

    def uia_tools_test():
        pwsh = shutil.which("pwsh")
        if not pwsh:
            raise RuntimeError("pwsh 不可用")
        for script in (mouse, uia_query, resume_capture):
            if not script.is_file():
                raise RuntimeError(f"缺少脚本：{script.name}")
            command = (
                "$tokens=$null;$errors=$null;"
                f"[void][System.Management.Automation.Language.Parser]::ParseFile('{str(script)}',[ref]$tokens,[ref]$errors);"
                "if($errors.Count){$errors | ForEach-Object {$_.Message}; exit 1}"
            )
            result = subprocess.run([pwsh, "-NoProfile", "-Command", command],
                                    capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
            if result.returncode:
                raise RuntimeError((result.stderr or result.stdout or "PowerShell parse failed")[-500:])
        return "UIA 查询、受控点击与简历采集脚本语法通过"

    checks.append(check("UIA 查询与点击脚本", uia_tools_test))
    checks.append(check("Codex 最大化", lambda: __import__("progress_window").ensure_codex_maximized()[1] or (_ for _ in ()).throw(RuntimeError("Codex 窗口未最大化"))))

    def idle_mouse_test():
        result = subprocess.run([sys.executable, str(idle_mouse), "--x", "120", "--y", "120", "--radius", "20",
                                 "--duration-ms", "20000", "--active-ms", "5000", "--rest-ms", "5000",
                                 "--content-bounds", "0,0,240,240", "--dry-run"],
                                capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
        if result.returncode:
            raise RuntimeError((result.stderr or result.stdout or "idle mouse dry run failed")[-500:])
        data = json.loads(result.stdout)
        if data.get("state") != "dry_run" or data.get("points", 0) < 10 or data.get("activePulses") != 2:
            raise RuntimeError("等待鼠标 dry run 结果异常")
        return "20 秒等待生成 2 个 5 秒移动脉冲，未移动鼠标"

    checks.append(check("等待鼠标轨迹 dry run", idle_mouse_test))
    checks.append({"name": "内置浏览器坐标与 Win32 句柄", "status": "通过" if args.browser_verified else "待核验",
                   "detail": args.browser_evidence.strip() if args.browser_verified else "需在本机现场只读校准；opaque browser ID 不等于 HWND"})
    passed = all(item["status"] == "通过" for item in checks)
    cn = timezone(timedelta(hours=8))
    stamp = datetime.now(cn)
    host = socket.gethostname().replace("/", "_").replace("\\", "_")
    report = folder / f"{host}-{stamp:%Y%m%d-%H%M%S}-标准化测试.md"
    lines = ["# BOSS 招聘 Skill 标准化测试", "", f"- 电脑：{host}", f"- 时间：{stamp.isoformat(timespec='seconds')}",
             f"- 系统：{platform.platform()}", f"- Python：{sys.executable}", f"- 结果：{'通过' if passed else '未通过'}", "",
             "| 检查项 | 状态 | 详情 |", "|---|---|---|"]
    for item in checks:
        lines.append(f"| {item['name']} | {item['status']} | {item['detail'].replace('|', '/').replace(chr(10), ' ')} |")
    lines += ["", "未执行真实发送；本报告不代表平台授权或发送成功率。", ""]
    folder.mkdir(parents=True, exist_ok=True)
    report.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"state": "passed" if passed else "not_ready", "report": str(report), "checks": checks}, ensure_ascii=False))


if __name__ == "__main__":
    main()
