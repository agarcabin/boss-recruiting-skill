"""Integration checks for score gate, workbook, status sync and resume/PDF archives."""

import base64
import contextlib
import io
import json
import tempfile
import unittest
from argparse import Namespace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from reportlab.pdfgen import canvas

import candidate_log
import ledger
import report_paths
from test_ledger import PROFILE


def invoke(fn, db, **kwargs):
    stream = io.StringIO()
    with contextlib.redirect_stdout(stream):
        fn(db, Namespace(**kwargs))
    return json.loads(stream.getvalue())


class CandidateLogTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        root_override = patch.object(ledger, "PROFILE_ROOT", self.root)
        root_override.start()
        self.addCleanup(root_override.stop)
        self.profile = self.root / "profile.md"
        self.profile.write_text(PROFILE, encoding="utf-8")
        self.db = ledger.connect(self.root / "ledger.db")
        candidate_log.table(self.db)
        self.run = invoke(ledger.start, self.db, profile=str(self.profile), target=5, duration_min=60, min_score=70)["run_id"]
        self.candidate = "https://example.test/candidate/one"
        self.resume_image = self.root / "resume-long.png"
        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        self.resume_image.write_bytes(png + b"\0" * 200)
        self.resume_text = self.root / "resume.ocr.txt"
        self.resume_text.write_text("简历 OCR 文本" * 20, encoding="utf-8")
        self.resume_json = self.root / "resume.ocr.json"
        self.resume_json.write_text('{"state":"ok","lineCount":1}', encoding="utf-8")

    def tearDown(self):
        self.db.close()

    def assessment(self, score=75, enabled=True, decision="greet", hard_requirements=True):
        stamp = datetime.now(timezone.utc).isoformat()
        payload = {"candidate_id": self.candidate, "candidate_name": "测试候选人", "started_at": stamp,
                   "coworker_progress": {"state": "empty", "evidence": "Test UIA coworker panel explicitly empty"},
                   "electronic_resume_enabled": enabled,
                   "overall_score": score, "hard_requirements": hard_requirements,
                   "decision": decision, "pros": "客户开发", "cons": "无额外业绩信息"}
        if enabled:
            payload["electronic"] = {"score": score, "evidence": "电子简历显示海外销售及客户开发"}
            payload["resume_evidence"] = {
                "image": str(self.resume_image),
                "ocr_text": str(self.resume_text),
                "ocr_json": str(self.resume_json),
            }
        else:
            payload["list"] = {"score": score, "evidence": "列表显示海外销售及客户开发"}
        return payload

    def record(self, payload):
        begun = invoke(ledger.candidate_begin, self.db, run=self.run, candidate_id=payload["candidate_id"])
        payload["started_at"] = begun["candidate_started_at"]
        path = self.root / "input.json"
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        return invoke(candidate_log.record, self.db, run=self.run, input=path, output_root=self.root)

    def test_score_gate_and_daily_workbook(self):
        with self.assertRaisesRegex(ValueError, "score >= 70"):
            self.record(self.assessment(69))
        result = self.record(self.assessment(75))
        self.assertEqual(Path(result["report"]).name, "Test Company-Sales-投递记录.xlsx")
        workbook = load_workbook(result["report"])
        try:
            summary = workbook["候选人简报"]
            headers = [cell.value for cell in summary[1]]
            values = [cell.value for cell in summary[2]]
            self.assertEqual(headers, candidate_log.SUMMARY_HEADERS)
            self.assertEqual(headers[0], "录入时间")
            self.assertEqual(headers[-1], "查看简历")
            self.assertEqual(values[headers.index("候选人")], "测试候选人")
            self.assertEqual(values[headers.index("综合分")], 75)
            self.assertEqual(values[headers.index("判定")], "打招呼")
            self.assertEqual(values[headers.index("判断时长")], "待完成")
            self.assertNotIn("项目", headers)
            self.assertNotIn("岗位", headers)
            self.assertNotIn("开始时间", headers)
            link = summary.cell(row=2, column=len(headers)).hyperlink
            self.assertTrue(link)
            self.assertTrue(link.target.startswith("简历图片/"))
            self.assertTrue((Path(result["report"]).parent / link.target).is_file())
            for cell in summary[2]:
                self.assertEqual(cell.fill.fgColor.rgb[-6:], "C6EFCE")
            details = workbook["详细信息"]
            self.assertEqual(details.sheet_state, "hidden")
            detail_headers = [cell.value for cell in details[1]]
            detail_values = [cell.value for cell in details[2]]
            self.assertEqual(detail_values[detail_headers.index("项目")], "Test Company")
            self.assertEqual(detail_values[detail_headers.index("岗位")], "Sales")
            self.assertEqual(detail_values[detail_headers.index("判断时长")], "待完成")
            self.assertIsNone(detail_values[detail_headers.index("判断时长（秒）")])
            self.assertEqual(detail_values[detail_headers.index("简历归档")], "saved")
        finally:
            workbook.close()
        self.assertEqual(invoke(ledger.reserve, self.db, run=self.run, candidate_id=self.candidate)["state"], "ready")

    def test_review_candidate_row_is_yellow(self):
        result = self.record(self.assessment(score=72, decision="review", hard_requirements=False))
        workbook = load_workbook(result["report"])
        try:
            for cell in workbook["候选人简报"][2]:
                self.assertEqual(cell.fill.fgColor.rgb[-6:], "FFEB9C")
        finally:
            workbook.close()

    def test_electronic_resume_score_is_required_when_enabled(self):
        payload = self.assessment()
        payload.pop("electronic")
        with self.assertRaisesRegex(ValueError, "electronic assessment required"):
            self.record(payload)

    def test_send_pdf_archive_blocks_next_candidate_until_complete(self):
        self.record(self.assessment())
        held = invoke(ledger.reserve, self.db, run=self.run, candidate_id=self.candidate)
        invoke(ledger.resolve, self.db, token=held["token"], result="sent", min_gap=64, max_gap=64)
        next_candidate = "https://example.test/candidate/two"
        self.assertEqual(invoke(ledger.reserve, self.db, run=self.run, candidate_id=next_candidate)["reason"], "report_not_synced")
        invoke(candidate_log.sync, self.db, run=self.run, candidate_id=self.candidate, output_root=self.root)
        self.assertEqual(invoke(ledger.reserve, self.db, run=self.run, candidate_id=next_candidate)["reason"], "pdf_archive_pending")

        pdf = self.root / "printed.pdf"
        document = canvas.Canvas(str(pdf))
        document.drawString(72, 700, "resume")
        document.save()
        archived = invoke(candidate_log.archive, self.db, run=self.run, candidate_id=self.candidate, pdf=str(pdf), output_root=self.root)
        self.assertTrue(Path(archived["pdf"]).is_file())
        self.assertEqual(invoke(ledger.candidate_finish, self.db, run=self.run, candidate_id=self.candidate)["reason"], "candidate_timing_pending")
        ended_at = ledger.iso(ledger.now())
        timing = invoke(candidate_log.finalize_timing, self.db, run=self.run, candidate_id=self.candidate,
                        ended_at=ended_at, output_root=self.root)
        self.assertEqual(timing["state"], "timing_finalized")
        self.assertEqual(invoke(ledger.candidate_finish, self.db, run=self.run, candidate_id=self.candidate)["state"], "finished_candidate")
        workbook = load_workbook(archived["report"], read_only=True)
        try:
            detail = workbook["详细信息"]
            headers = [cell.value for cell in detail[1]]
            values = [cell.value for cell in detail[2]]
            self.assertEqual(values[headers.index("发送状态")], "sent")
            self.assertEqual(values[headers.index("PDF归档")], "saved")
            self.assertEqual(values[headers.index("简历归档")], "saved")
            self.assertTrue(Path(values[headers.index("简历图片")]).is_file())
            self.assertEqual(values[headers.index("结束时间")], ended_at)
        finally:
            workbook.close()

    def test_timing_finalization_is_idempotent_and_uses_click_timestamp(self):
        result = self.record(self.assessment(score=72, decision="review", hard_requirements=False))
        invoke(ledger.note, self.db, run=self.run, candidate_id=self.candidate, result="review")
        invoke(candidate_log.sync, self.db, run=self.run, candidate_id=self.candidate, output_root=self.root)
        stamp = ledger.iso(ledger.now())
        finished = invoke(candidate_log.finalize_timing, self.db, run=self.run, candidate_id=self.candidate,
                          ended_at=stamp, output_root=self.root)
        self.assertEqual(finished["candidate_ended_at"], stamp)
        self.assertEqual(invoke(candidate_log.finalize_timing, self.db, run=self.run,
                                candidate_id=self.candidate, ended_at=stamp,
                                output_root=self.root)["state"], "timing_already_finalized")
        with self.assertRaisesRegex(ValueError, "candidate_timing_already_finalized"):
            invoke(candidate_log.finalize_timing, self.db, run=self.run, candidate_id=self.candidate,
                   ended_at=ledger.iso(ledger.now() + timedelta(seconds=1)), output_root=self.root)
        self.assertEqual(invoke(ledger.candidate_finish, self.db, run=self.run, candidate_id=self.candidate,
                                ended_at=ledger.iso(ledger.now() + timedelta(seconds=1)))["reason"], "candidate_end_mismatch")
        self.assertEqual(invoke(ledger.candidate_finish, self.db, run=self.run, candidate_id=self.candidate,
                                ended_at=stamp)["state"], "finished_candidate")
        workbook = load_workbook(result["report"], read_only=True)
        try:
            self.assertNotEqual(workbook["候选人简报"]["E2"].value, "待完成")
        finally:
            workbook.close()

    def test_not_found_policy_survives_record_and_reserve(self):
        payload = self.assessment(75)
        payload['coworker_progress'] = {'state': 'not_found', 'lookup_completed': True,
                                       'evidence': 'Directed checks completed; no coworker header in live frames'}
        self.record(payload)
        result = invoke(ledger.reserve, self.db, run=self.run, candidate_id=self.candidate)
        self.assertEqual(result['state'], 'ready')

    def test_fractional_action_time_can_be_finalized_twice_and_finished(self):
        self.record(self.assessment(score=55, decision="review", hard_requirements=False))
        invoke(ledger.note, self.db, run=self.run, candidate_id=self.candidate, result="review")
        invoke(candidate_log.sync, self.db, run=self.run, candidate_id=self.candidate, output_root=self.root)
        stamp = ledger.now().replace(microsecond=123456).isoformat()
        first = invoke(candidate_log.finalize_timing, self.db, run=self.run,
                       candidate_id=self.candidate, ended_at=stamp, output_root=self.root)
        repeated = invoke(candidate_log.finalize_timing, self.db, run=self.run,
                          candidate_id=self.candidate, ended_at=stamp, output_root=self.root)
        self.assertEqual(repeated["state"], "timing_already_finalized")
        finished = invoke(ledger.candidate_finish, self.db, run=self.run,
                          candidate_id=self.candidate, ended_at=stamp)
        self.assertEqual(finished["state"], "finished_candidate")
        self.assertEqual(first["candidate_ended_at"], finished["candidate_ended_at"])

    def test_reports_and_rebuild_separate_positions_on_the_same_day(self):
        first = self.record(self.assessment(score=72, decision="review", hard_requirements=False))
        profile = self.root / "assembly.md"
        profile.write_text(PROFILE.replace("position: Sales", "position: Assembly"), encoding="utf-8")
        self.run = invoke(ledger.start, self.db, profile=str(profile), target=5,
                          duration_min=60, min_score=70)["run_id"]
        self.candidate = "https://example.test/candidate/another-position"
        second = self.record(self.assessment(score=72, decision="review", hard_requirements=False))
        self.assertNotEqual(first["report"], second["report"])
        self.assertEqual(Path(second["report"]).name, "Test Company-Assembly-投递记录.xlsx")
        day = datetime.now(candidate_log.CN_TZ).date().isoformat()
        rebuilt = invoke(candidate_log.rebuild, self.db, day=day, output_file=None, output_root=self.root)
        self.assertEqual(set(rebuilt["reports"]), {first["report"], second["report"]})
        for result, position in [(first, "Sales"), (second, "Assembly")]:
            wb = load_workbook(result["report"])
            try:
                self.assertEqual(wb["候选人简报"].max_row, 2)
                details = wb["详细信息"]
                headers = [c.value for c in details[1]]
                self.assertEqual(details.cell(2, headers.index("岗位") + 1).value, position)
            finally:
                wb.close()

    def test_custom_report_names_keep_prefix_and_reject_traversal(self):
        self.assertEqual(report_paths.report_filename(self.profile, "另存副本.xlsx"),
                         "Test Company-Sales-另存副本.xlsx")
        self.assertEqual(report_paths.report_filename(self.profile, "Test Company-Sales-另存副本.xlsx"),
                         "Test Company-Sales-另存副本.xlsx")
        with self.assertRaisesRegex(ValueError, "output_file"):
            report_paths.report_filename(self.profile, "../outside.xlsx")
        a=self.root/"a.md";b=self.root/"b.md"
        a.write_text(PROFILE.replace("Test Company", "Test/Company"),encoding="utf-8")
        b.write_text(PROFILE.replace("Test Company", "Test\\Company"),encoding="utf-8")
        self.assertNotEqual(report_paths.report_filename(a),report_paths.report_filename(b))

    def test_local_name_dedup_is_off_even_when_matching_history_exists(self):
        self.record(self.assessment(score=72, decision="review", hard_requirements=False))
        for enabled, expected in [(False, False), (True, True)]:
            stream = io.StringIO()
            with contextlib.redirect_stdout(stream):
                candidate_log.lookup_name(Namespace(name="测试候选人", db=self.root / "ledger.db",
                                                    enable_local_excel_dedup=enabled))
            result = json.loads(stream.getvalue())
            self.assertEqual(result["duplicate"], expected)
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            candidate_log.lookup_name(Namespace(name="测试候选人", db=self.root / "does-not-exist.db"))
        self.assertEqual(json.loads(stream.getvalue())["state"], "disabled")

    def test_coworker_check_required_and_old_first_greeting_blocks_send(self):
        payload = self.assessment()
        payload.pop("coworker_progress")
        with self.assertRaisesRegex(ValueError, "coworker_check_required"):
            candidate_log.validate(payload)
        payload["coworker_progress"] = {
            "state": "first_record", "first_greeting_at": "2020-06-05 09:22",
            "evidence": "UIA first row; later rows must not override this one"}
        with self.assertRaisesRegex(ValueError, "first_coworker_older_than_7_days"):
            candidate_log.validate(payload)

    def test_reserve_rechecks_coworker_time_boundary(self):
        at = ledger.now()
        payload = self.assessment()
        payload["coworker_progress"] = {
            "state": "first_record", "first_greeting_at": (at-timedelta(days=7)+timedelta(seconds=60)).isoformat(),
            "evidence": "UIA first coworker row"}
        self.record(payload)
        with patch.object(ledger, "now", return_value=at+timedelta(seconds=61)):
            result = invoke(ledger.reserve, self.db, run=self.run, candidate_id=self.candidate)
        self.assertEqual(result["reason"], "first_coworker_older_than_7_days")

    def test_reserve_blocks_legacy_unverified_coworker_state(self):
        self.record(self.assessment())
        self.db.execute("UPDATE assessments SET coworker_progress=NULL")
        result = invoke(ledger.reserve, self.db, run=self.run, candidate_id=self.candidate)
        self.assertEqual(result["reason"], "coworker_check_required")

    def test_old_coworker_skip_is_reported_unscored_and_can_finish(self):
        invoke(ledger.candidate_begin, self.db, run=self.run, candidate_id=self.candidate)
        path = self.root / "coworker.json"
        path.write_text(json.dumps({"state":"first_record", "first_greeting_at":"2020-06-05 09:22",
                                    "evidence":"UIA first visible coworker row"}), encoding="utf-8")
        result = invoke(candidate_log.record_coworker_skip, self.db, run=self.run,
                        candidate_id=self.candidate, candidate_name="测试候选人", input=path, output_root=self.root)
        wb = load_workbook(result["report"])
        try:
            self.assertIsNone(wb["候选人简报"]["C2"].value)
            self.assertEqual(wb["候选人简报"]["D2"].value, "重复跳过")
        finally:
            wb.close()
        invoke(ledger.note, self.db, run=self.run, candidate_id=self.candidate, result="skipped")
        invoke(candidate_log.sync, self.db, run=self.run, candidate_id=self.candidate, output_root=self.root)
        stamp=ledger.iso(ledger.now())
        invoke(candidate_log.finalize_timing, self.db, run=self.run, candidate_id=self.candidate,
               ended_at=stamp, output_root=self.root)
        result=invoke(ledger.candidate_finish, self.db, run=self.run, candidate_id=self.candidate, ended_at=stamp)
        self.assertEqual(result["state"], "finished_candidate")


if __name__ == "__main__":
    unittest.main()
