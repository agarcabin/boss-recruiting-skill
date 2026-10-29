"""Behavioral coverage of the remaining October 29 update plan; no BOSS sends."""
import contextlib
import io
import json
import tempfile
import unittest
from datetime import datetime,timedelta,timezone
from pathlib import Path
from unittest.mock import Mock,patch
from argparse import Namespace
import candidate_checks
import candidate_log
import capture_resume
import daily_metrics
import ledger
import progress_window
import quick_start
import recruiting_rules
import ui_features
from keep_awake import KeepAwake
import test_relaxed_controls as controls
call=controls.call


class RulesTests(unittest.TestCase):
    def setUp(self):
        self.rules=json.loads(recruiting_rules.CONFIG.read_text(encoding='utf-8'))
        self.rules.update(policy='meikang',mode='resume')

    def assess(self,text='',**kw):
        payload=dict(overall_score=40,hard_requirements=False,keyword_text=text)
        payload.update(kw)
        return recruiting_rules.evaluate(self.rules,payload,70)

    def test_all_28_employers_and_each_business_word(self):
        self.assertEqual(len(self.rules['meikang_companies']),28)
        for company in self.rules['meikang_companies']:
            for word in ('海外','出口','销售','业务'):
                result=self.assess(work_experience=[dict(company=company+'股份有限公司',duties='负责'+word)])
                self.assertTrue(result['eligible'],(company,word))
                self.assertIn('company_business',result['direct_channels'])

    def test_employer_and_business_can_be_separate_jobs(self):
        result=self.assess(work_experience=[dict(company='迈瑞医疗',duties='研发'),dict(company='其他公司',duties='销售经理')])
        self.assertTrue(result['eligible'])

    def test_customer_reference_and_negation_do_not_trigger(self):
        samples=['工作经历\n其他公司\n客户：迈瑞医疗，负责销售','就职于其他公司，客户名单包括迈瑞医疗，负责销售',
                 '工作经历\n迈瑞医疗\n没有海外销售业务经验','求职意向：想在迈瑞医疗做销售']
        for text in samples:
            self.assertFalse(self.assess(text)['eligible'],text)

    def test_plain_work_section_and_explicit_employment_trigger(self):
        self.assertTrue(self.assess('工作经历\n迈瑞医疗股份有限公司\n2019.01-2023.01\n销售经理\n教育经历\n大学')['eligible'])
        self.assertTrue(self.assess('曾在新产业生物任职销售经理')['eligible'])

    def test_product_export_requires_own_work_and_overseas(self):
        for product in ('生化','化学发光','质谱'):
            self.assertTrue(self.assess(work_experience=[dict(company='其他公司',duties=f'负责{product}试剂出口，开拓欧洲市场')])['eligible'])
        self.assertFalse(self.assess('海外留学期间了解质谱产品出口')['eligible'])
        self.assertFalse(self.assess(work_experience=[dict(company='其他公司',duties='负责质谱研发，国内市场')])['eligible'])
        self.assertFalse(self.assess(work_experience=[dict(company='其他公司',duties='没有质谱出口经验，曾出差海外')])['eligible'])

    def test_unrelated_export_clause_does_not_trigger(self):
        result=self.assess(work_experience=[dict(company='其他公司',duties='负责生化仪器研发。服装出口海外。')])
        self.assertFalse(result['eligible'])

    def test_experience_only_deducts_and_direct_still_passes(self):
        for years,deduct in ((3,0),(2,5),(1,10),(0,10)):
            result=self.assess(work_experience=[dict(company='迈瑞医疗',duties='销售')],related_experience_years=years)
            self.assertTrue(result['eligible'])
            self.assertEqual(result['experience_penalty'],deduct)
            self.assertEqual(result['score'],40-deduct)
        self.assertEqual(self.assess(related_experience_years=None)['experience_penalty'],0)

    def test_years_failed_is_not_hard_exclusion_and_penalty_not_applied_twice(self):
        p=dict(overall_score=90,related_experience_years=2,hard_requirements=False,requirement_states=dict(years='failed',education='met'),keyword_text='医疗器械海外销售')
        first=recruiting_rules.evaluate(self.rules,p,70)
        self.assertTrue(first['eligible'])
        p.update(overall_score=first['score'],score_before_experience=first['score_before_experience'])
        self.assertEqual(recruiting_rules.evaluate(self.rules,p,70)['score'],first['score'])

    def test_worker_colleague_check_precedes_role_keywords(self):
        rules=dict(self.rules,policy='worker',mode='card')
        for state in ('present','unknown'):
            result=recruiting_rules.evaluate(rules,dict(card=dict(advantage='组装',expectation='普工',coworker_contact=dict(state=state,lookup_completed=True,evidence='当前卡片检查'))),70)
            self.assertFalse(result['eligible'])
            self.assertEqual(result['hits'],[])
        result=recruiting_rules.evaluate(rules,dict(card=dict(advantage='组装',coworker_contact=dict(state='absent',lookup_completed=True,evidence='已查当前卡片'))),70)
        self.assertTrue(result['eligible'])

    def test_company_name_fragments_and_explicit_business_negation(self):
        for name in ('深圳迈瑞生物医疗电子股份有限公司','深圳开立生物医疗科技股份有限公司','深圳市理邦精密仪器股份有限公司'):
            self.assertTrue(self.assess(work_experience=[dict(company=name,duties='负责销售')])['eligible'])
        self.assertFalse(self.assess(work_experience=[dict(company='迈瑞医疗',duties='不负责销售，不涉及海外业务')])['eligible'])


class FlowTests(controls.RelaxedControlsTests):
    # Inherit the 1.0.12 regression suite so removed controls stay removed.
    def test_optional_capture_does_not_require_active_candidate(self):
        candidate_checks.record(self.db,self.run,'not-registered',dict(state='empty',evidence='当前同事面板明确为空'))
        call(ledger.candidate_finish,self.db,run=self.run,candidate_id='not-registered')
        self.assertEqual(capture_resume.check_candidate(self.db,self.run,'not-registered',-999)['run_id'],self.run)

    def test_missing_or_skip_colleague_prevents_capture_and_rule_analysis(self):
        with self.assertRaisesRegex(ValueError,'coworker_check_before_analysis'):
            capture_resume.check_candidate(self.db,self.run,'candidate-one',0)
        old=dict(state='first_record',first_greeting_at=ledger.iso(ledger.now()-timedelta(days=9)),evidence='UIA首条同事日期')
        with patch.object(recruiting_rules,'evaluate',side_effect=AssertionError('must not analyze')):
            with self.assertRaisesRegex(ValueError,'coworker_check_before_analysis'):
                self.record(coworker_progress=old)

    def test_colleague_result_is_not_reused_for_next_visit(self):
        candidate_checks.record(self.db,self.run,'candidate-one',dict(state='empty',evidence='首次实际检查'))
        call(ledger.candidate_finish,self.db,run=self.run,candidate_id='candidate-one')
        call(ledger.candidate_begin,self.db,run=self.run,candidate_id='candidate-one')
        with self.assertRaisesRegex(ValueError,'coworker_check_before_analysis'):
            candidate_checks.require_resume(self.db,self.run,'candidate-one')

    def test_green_rows_only_after_confirmed_success(self):
        self.record(); path=Path(self.record('candidate-two')['report'])
        self.assertEqual(daily_metrics.workbook_counts(path),dict(viewed=2,sent=0))
        result=self.reserve('candidate-two');self.resolve(result['token'])
        call(candidate_log.sync,self.db,run=self.run,candidate_id='candidate-two',output_root=self.root)
        self.assertEqual(daily_metrics.workbook_counts(path),dict(viewed=2,sent=1))
        current=daily_metrics.daily_counts(self.db,ledger.run_row(self.db,self.run),self.root)
        self.assertEqual((current['daily_viewed'],current['daily_sent']),(2,1))

    def test_daily_counts_continue_across_runs_but_target_is_per_run(self):
        self.record();self.resolve(self.reserve()['token'])
        call(candidate_log.sync,self.db,run=self.run,candidate_id='candidate-one',output_root=self.root)
        self.run=self.start(target=50)['run_id'];self.record('candidate-two')
        current=progress_window.snapshot(self.db,self.run,self.root)
        self.assertEqual((current['daily_viewed'],current['daily_sent'],current['sent']),(2,1,0))

    def test_daily_counts_cross_midnight_and_other_jobs_are_separate(self):
        self.record()
        tomorrow=ledger.now()+timedelta(days=1)
        self.assertEqual(daily_metrics.projected_counts(self.db,self.profile,tomorrow)['viewed'],0)
        other=self.root/'other.md';other.write_text(self.profile.read_text(encoding='utf-8').replace('position: Sales','position: Other'),encoding='utf-8')
        self.assertEqual(daily_metrics.projected_counts(self.db,other,ledger.now())['viewed'],0)

    def test_locked_report_shows_actual_counts_and_pending(self):
        first=self.record();path=Path(first['report'])
        with patch.object(candidate_log,'workbook',side_effect=PermissionError('locked')):
            self.record('candidate-two')
        current=daily_metrics.daily_counts(self.db,ledger.run_row(self.db,self.run),self.root)
        self.assertEqual(current['daily_viewed'],1)
        self.assertTrue(current['report_pending'])

    def test_favorite_review_unsaved_is_clicked_once_and_verified(self):
        self.record(decision='review')
        reader=Mock(side_effect=[dict(state='unsaved',target=dict(name='收藏')),dict(state='saved')]);click=Mock()
        result=ui_features.favorite(self.db,self.run,'candidate-one',reader=reader,click=click)
        self.assertEqual(result['state'],'saved');click.assert_called_once()

    def test_favorite_sent_already_saved_does_not_toggle(self):
        self.record();self.resolve(self.reserve()['token']);click=Mock()
        result=ui_features.favorite(self.db,self.run,'candidate-one',reader=Mock(return_value=dict(state='saved')),click=click)
        self.assertEqual(result['state'],'already_saved');click.assert_not_called()

    def test_unknown_send_and_skip_do_not_favorite(self):
        self.record();self.resolve(self.reserve()['token'],'unknown');reader=Mock()
        self.assertEqual(ui_features.favorite(self.db,self.run,'candidate-one',reader=reader)['state'],'not_applicable')
        reader.assert_not_called()

    def test_favorite_unknown_after_click_does_not_retry(self):
        self.record(decision='review');click=Mock();reader=Mock(side_effect=[dict(state='unsaved',target={}),dict(state='unavailable')])
        self.assertEqual(ui_features.favorite(self.db,self.run,'candidate-one',reader=reader,click=click)['state'],'pending')
        click.assert_called_once()

    def test_direct_channel_with_low_score_and_two_years_passes_record_and_reserve(self):
        rules=json.loads(recruiting_rules.CONFIG.read_text(encoding='utf-8'));rules.update(policy='meikang',mode='resume',greeting_mode='single')
        self.db.execute('UPDATE runs SET rule_snapshot=? WHERE run_id=?',(json.dumps(rules),self.run))
        result=self.record(hard_requirements=False,overall_score=50,electronic=dict(score=50,evidence='实际迈瑞医疗任职销售经历'),
                           work_experience=[dict(company='迈瑞医疗',duties='负责销售业务')],related_experience_years=2)
        self.assertEqual(result['decision'],'greet')
        self.assertEqual(self.db.execute('SELECT overall_score FROM assessments').fetchone()[0],45)
        self.assertEqual(self.reserve()['state'],'ready')

    def test_colleague_skip_and_unknown_are_logged_without_analysis_or_score(self):
        observations=[dict(state='first_record',first_greeting_at=ledger.iso(ledger.now()-timedelta(days=9)),evidence='当前同事首条日期'),dict(state='unavailable',evidence='当前面板时间不可读')]
        for index,observation in enumerate(observations):
            with patch.object(recruiting_rules,'evaluate',side_effect=AssertionError('must not analyze')):
                result=self.record('candidate-check-'+str(index),decision='skip' if index==0 else 'review',coworker_progress=observation)
            self.assertEqual(result['state'],'recorded')
        from openpyxl import load_workbook
        wb=load_workbook(result['report'])
        try:
            self.assertIsNone(wb['候选人简报']['C2'].value)
            self.assertIsNone(wb['候选人简报']['C3'].value)
            self.assertEqual(wb['候选人简报']['D3'].value,'同事检查待处理')
        finally:wb.close()

    def test_favorite_result_is_exported_in_hidden_evidence_sheet(self):
        self.record(decision='review')
        ui_features.favorite(self.db,self.run,'candidate-one',reader=Mock(return_value=dict(state='saved')),click=Mock())
        path=candidate_log.workbook(self.db,self.root,ledger.iso(ledger.now()),profile_path=self.profile)
        from openpyxl import load_workbook
        wb=load_workbook(path)
        try:self.assertEqual(wb['评估证据']['I2'].value,'already_saved')
        finally:wb.close()

    def test_favorite_read_failure_is_recorded_without_blocking(self):
        self.record(decision='review')
        result=ui_features.favorite(self.db,self.run,'candidate-one',reader=Mock(side_effect=RuntimeError('UIA unavailable')))
        self.assertEqual(result['state'],'pending')
        self.assertEqual(self.db.execute('SELECT state FROM favorite_results').fetchone()[0],'pending')
        self.record('candidate-two')
        self.assertEqual(self.reserve('candidate-two')['state'],'ready')


class StartupAndUITests(unittest.TestCase):
    def setUp(self):
        self.projects=[dict(profile='meikang.md',company='美康生物',position='国际销售经理',approval='pending')]

    def test_quick_start_targets_title_and_filter_pause(self):
        at=datetime(2026,10,29,17,tzinfo=timezone.utc)
        result=quick_start.parse('帮我打招呼，打美康 50 人',self.projects,at)
        self.assertEqual((result['state'],result['target'],result['max_views']),('await_filters',50,300))
        self.assertTrue(result['chat_title'].endswith('2026-10-30'))
        self.assertTrue(result['approval_warning'])

    def test_view_target_and_unlimited_override(self):
        result=quick_start.parse('帮美康招呼500人，查看1200次',self.projects)
        self.assertEqual((result['target'],result['max_views']),(500,1200))
        self.assertEqual(quick_start.parse('帮美康看100人',self.projects)['max_views'],100)

    def test_ambiguous_job_requires_only_job_selection(self):
        projects=self.projects+[dict(profile='other.md',company='美康生物',position='其他岗位')]
        self.assertEqual(quick_start.parse('打美康50人',projects)['state'],'needs_selection')

    def test_login_popup_and_sidebar_are_idempotent(self):
        click=Mock();reader=Mock(side_effect=[dict(state='present',target={}),dict(state='absent')])
        self.assertEqual(ui_features.once('Login',reader=reader,click=click)['state'],'confirmed');click.assert_called_once()
        click.reset_mock()
        self.assertEqual(ui_features.once('Sidebar',reader=Mock(return_value=dict(state='collapsed')),click=click)['state'],'collapsed');click.assert_not_called()

    def test_awake_running_pause_resume_and_release(self):
        apply=Mock(return_value=1);wake=KeepAwake(apply)
        for state in (True,True,False,True,False):wake.set_running(state)
        self.assertEqual([call.args[0] for call in apply.call_args_list],[0x80000003,0x80000000,0x80000003,0x80000000])

    def test_button_centering_and_no_overlap_for_three_dpi_scales(self):
        for scale in (1,1.25,1.5):
            g=progress_window.button_geometry(scale)
            self.assertLessEqual(abs(2*g['y']+g['height']-round(progress_window.PANEL_HEIGHT/scale)),1)
            self.assertLessEqual(2*g['x']+sum(g['widths'])+2*g['gap'],round(progress_window.BUTTON_WIDTH/scale)+2)

    def test_eta_uses_both_user_targets_and_actual_success_rate(self):
        run=dict(target=50,max_views=300)
        minutes,basis=daily_metrics.estimate(run,100,25,6,600)
        self.assertEqual((minutes,basis),(10,'招呼目标'))
        self.assertEqual(daily_metrics.estimate(run,0,0,None,0),(None,'估算中'))

    def test_visual_card_check_parses_role_only_after_no_colleague_icon(self):
        import card_observation
        with tempfile.TemporaryDirectory() as temporary:
            image=Path(temporary)/'card.png';image.write_bytes(b'fixture')
            observation=dict(raw_text='测试人 优势：组装 期望：普工',card={})
            skipped=card_observation.complete(observation,'present','在本地截图看见同事标记',image)
            self.assertEqual(skipped['card']['advantage'],'')
            clear=card_observation.complete(observation,'absent','实际卡片截图无同事标记',image)
            self.assertEqual(clear['card']['expectation'],'普工')


if __name__=='__main__':unittest.main()
